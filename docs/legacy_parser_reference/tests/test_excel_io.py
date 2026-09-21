"""M1b Excel 导入：模板、文本学号、引用、重复、缺列、人数、公式、定位。"""

import io

from openpyxl import Workbook, load_workbook

from excel_io import (
    SHEET_COURSES,
    SHEET_OVERRIDES,
    SHEET_TASKS,
    SHEET_VOLUNTEERS,
    create_template_workbook,
    read_workbook,
)


def _make_workbook(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=name)
        first = False
        ws.title = name
        for r in rows:
            ws.append(list(r))
    # 文本列：学号等按文本存储，避免前导零丢失。
    if SHEET_VOLUNTEERS in wb.sheetnames:
        ws = wb[SHEET_VOLUNTEERS]
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=1)
            if c.value is not None:
                c.value = str(c.value)
                c.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valid_sheets(**kw) -> dict[str, list[list[object]]]:
    base: dict[str, list[list[object]]] = {
        SHEET_VOLUNTEERS: [
            ["学号", "姓名", "班级"],
            ["V001", "张三", "A"],
            ["V002", "李四", "B"],
        ],
        SHEET_COURSES: [
            ["课程编号", "班级", "课程名称", "周次", "星期", "节次"],
            ["C001", "A", "高数", "1-8", "周一", "1-2"],
            ["C002", "A", "英语", "1-8", "周一", "1-2"],
        ],
        SHEET_OVERRIDES: [
            ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ],
        SHEET_TASKS: [
            ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
            ["T001", "B", "3", "周一", "1-2", "101", 2],
        ],
    }
    base.update(kw)
    return base


def test_template_readable():
    data = create_template_workbook()
    snap, issues = read_workbook(data, 16)
    assert snap is not None, issues
    assert issues == [] or all(i.code == "DUPLICATE_OCCUPANCY" for i in issues)
    assert {v.student_no for v in snap.volunteers} >= {"V001", "V002"}
    assert len(snap.courses) > 0
    assert len(snap.tasks) > 0


def test_text_student_no_keeps_leading_zero():
    sheets = _valid_sheets()
    sheets[SHEET_VOLUNTEERS] = [
        ["student_no", "name", "class_id"],
        ["001", "张三", "A"],
    ]
    sheets[SHEET_COURSES] = [
        ["course_id", "class_id", "course_name", "weeks", "weekday", "period"],
        ["C001", "A", "高数", "1", "周一", "1-2"],
    ]
    sheets[SHEET_OVERRIDES] = [
        ["override_id", "student_no", "action", "course_id", "weeks", "weekday", "period", "remark"],
    ]
    sheets[SHEET_TASKS] = [
        ["task_id", "target_class_id", "weeks", "weekday", "period", "location", "required_people"],
        ["T001", "A", "2", "周二", "1-2", "", 1],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is not None, issues
    assert snap.volunteers[0].student_no == "001"
    assert snap.volunteers[0].id == "001"


def test_unknown_refs_block():
    # 未知学号
    sheets = _valid_sheets()
    sheets[SHEET_OVERRIDES] = [
        ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ["O001", "VX", "ADD", "", "1", "周一", "1-2", ""],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "UNKNOWN_VOLUNTEER" and i.sheet == "个人修正" and i.row == 2 for i in issues)

    # 取消未知课程
    sheets = _valid_sheets()
    sheets[SHEET_OVERRIDES] = [
        ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ["O001", "V001", "CANCEL_COURSE", "CX", "1", "周一", "1-2", ""],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "UNKNOWN_COURSE" for i in issues)

    # 任务目标班级未知
    sheets = _valid_sheets()
    sheets[SHEET_TASKS] = [
        ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
        ["T001", "ZZZ", "1", "周一", "1-2", "", 1],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "UNKNOWN_CLASS" and i.field == "target_class_id" for i in issues)


def test_duplicate_ids_block():
    sheets = _valid_sheets()
    sheets[SHEET_VOLUNTEERS] = [
        ["学号", "姓名", "班级"],
        ["V001", "张三", "A"],
        ["V001", "李四", "B"],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "DUPLICATE_ID" and i.field == "student_no" for i in issues)

    sheets = _valid_sheets()
    sheets[SHEET_OVERRIDES] = [
        ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ["O001", "V001", "ADD", "", "1", "周一", "1-2", ""],
        ["O001", "V002", "ADD", "", "1", "周一", "1-2", ""],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "DUPLICATE_ID" and i.sheet == "个人修正" for i in issues)

    sheets = _valid_sheets()
    sheets[SHEET_TASKS] = [
        ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
        ["T001", "B", "1", "周一", "1-2", "", 1],
        ["T001", "B", "2", "周一", "1-2", "", 1],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "DUPLICATE_ID" and i.sheet == "查课任务" for i in issues)


def test_missing_column_and_bad_people():
    sheets = _valid_sheets()
    sheets[SHEET_VOLUNTEERS] = [
        ["学号", "姓名"],  # 缺班级列
        ["V001", "张三"],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "MISSING_COLUMN" and i.field == "class_id" for i in issues)

    for bad in (0, -1, "abc", "2.5"):
        sheets = _valid_sheets()
        sheets[SHEET_TASKS] = [
            ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
            ["T001", "B", "1", "周一", "1-2", "", bad],
        ]
        snap, issues = read_workbook(_make_workbook(sheets), 16)
        assert snap is None, bad
        assert any(i.code in ("INVALID_REQUIRED_PEOPLE", "EMPTY_FIELD") for i in issues), bad


def test_formula_rejected_with_location():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_VOLUNTEERS
    ws.append(["学号", "姓名", "班级"])
    ws.append(["V001", "张三", "A"])
    ws.append(["V002", "李四", "B"])
    for name, header in [
        (SHEET_COURSES, ["课程编号", "班级", "课程名称", "周次", "星期", "节次"]),
        (SHEET_OVERRIDES, ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"]),
        (SHEET_TASKS, ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"]),
    ]:
        ws2 = wb.create_sheet(title=name)
        ws2.append(header)
    ws2 = wb[SHEET_COURSES]
    ws2.append(["C001", "A", "高数", "1-8", "周一", "1-2"])
    ws2 = wb[SHEET_TASKS]
    ws2.append(["T001", "B", "1", "周一", "1-2", "", 1])
    # 公式单元格：需要人数写成公式
    ws2.cell(row=2, column=7).value = "=1+1"
    assert ws2.cell(row=2, column=7).data_type == "f"
    buf = io.BytesIO()
    wb.save(buf)
    snap, issues = read_workbook(buf.getvalue(), 16)
    assert snap is None
    assert any(
        i.code == "FORMULA_CELL" and i.sheet == "查课任务" and i.row == 2 and i.field == "required_people"
        for i in issues
    )


def test_task_multi_slot_and_cancel_not_own():
    sheets = _valid_sheets()
    sheets[SHEET_TASKS] = [
        ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
        ["T001", "B", "1", "周一", "1-6", "", 1],  # 映射到 [1,2,3]
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "TASK_MULTIPLE_SLOTS" for i in issues)

    sheets = _valid_sheets()
    # V001 在 A 班，却取消 B 班课程（C003 属于 B）
    sheets[SHEET_COURSES] = [
        ["课程编号", "班级", "课程名称", "周次", "星期", "节次"],
        ["C001", "A", "高数", "1-8", "周一", "1-2"],
        ["C003", "B", "物理", "1-8", "周二", "1-2"],
    ]
    sheets[SHEET_OVERRIDES] = [
        ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ["O001", "V001", "CANCEL_COURSE", "C003", "1", "周二", "1-2", ""],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is None
    assert any(i.code == "CANCEL_NOT_OWN_COURSE" for i in issues)


def test_duplicate_add_same_slot_warns_not_blocks():
    sheets = _valid_sheets()
    sheets[SHEET_OVERRIDES] = [
        ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
        ["O001", "V001", "ADD", "", "2", "周二", "3-4", ""],
        ["O002", "V001", "ADD", "", "2", "周二", "3-4", ""],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is not None, issues
    assert any(i.code == "DUPLICATE_OCCUPANCY" for i in issues)


def test_missing_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_VOLUNTEERS
    ws.append(["学号", "姓名", "班级"])
    buf = io.BytesIO()
    wb.save(buf)
    snap, issues = read_workbook(buf.getvalue(), 16)
    assert snap is None
    assert any(i.code == "MISSING_SHEET" for i in issues)


def test_weeks_expand_and_task_id_stable():
    sheets = _valid_sheets()
    sheets[SHEET_TASKS] = [
        ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
        ["T009", "B", "1-2", "周一", "1-2", "101", 1],
    ]
    snap, issues = read_workbook(_make_workbook(sheets), 16)
    assert snap is not None, issues
    ids = sorted(t.id for t in snap.tasks)
    assert ids == ["T009@w01d1p1", "T009@w02d1p1"]


def test_excel_cancel_semantics_end_to_end():
    """内存工作簿：V001/V002 同班、两门同格课程，取消语义经读取+占用计算验证。"""
    from availability import build_occupied

    def _sheets_with_overrides(overrides_rows):
        return {
            SHEET_VOLUNTEERS: [
                ["学号", "姓名", "班级"],
                ["V001", "张三", "A"],
                ["V002", "李四", "A"],
            ],
            SHEET_COURSES: [
                ["课程编号", "班级", "课程名称", "周次", "星期", "节次"],
                ["C01", "A", "高数", "3", "周一", "1-2"],
                ["C02", "A", "英语", "3", "周一", "1-2"],
            ],
            SHEET_OVERRIDES: [
                ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
                *overrides_rows,
            ],
            SHEET_TASKS: [
                ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
                ["T001", "A", "3", "周二", "1-2", "", 1],
            ],
        }

    slot = (3, 1, 1)
    # 取消其中一门仍占用
    snap, issues = read_workbook(_make_workbook(_sheets_with_overrides([
        ["O01", "V001", "CANCEL_COURSE", "C01", "3", "周一", "1-2", ""],
    ])), 16)
    assert snap is not None, issues
    assert slot in build_occupied(snap)["V001"]

    # 取消两门后释放
    snap, issues = read_workbook(_make_workbook(_sheets_with_overrides([
        ["O01", "V001", "CANCEL_COURSE", "C01", "3", "周一", "1-2", ""],
        ["O02", "V001", "CANCEL_COURSE", "C02", "3", "周一", "1-2", ""],
    ])), 16)
    assert snap is not None, issues
    assert slot not in build_occupied(snap)["V001"]
    assert slot in build_occupied(snap)["V002"]

    # ADD 后重新占用
    snap, issues = read_workbook(_make_workbook(_sheets_with_overrides([
        ["O01", "V001", "CANCEL_COURSE", "C01", "3", "周一", "1-2", ""],
        ["O02", "V001", "CANCEL_COURSE", "C02", "3", "周一", "1-2", ""],
        ["O03", "V001", "ADD", "", "3", "周一", "1-2", ""],
    ])), 16)
    assert snap is not None, issues
    assert slot in build_occupied(snap)["V001"]
