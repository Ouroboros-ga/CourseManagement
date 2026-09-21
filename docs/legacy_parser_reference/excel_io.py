"""Excel 导入与模板（M1b）+ 结果导出（M2b）。

导入见 read_workbook；导出见 export_workbook（五表：总排班/个人排班/
未分配任务/冲突报告/次数统计）。用户文本一律按文本存储（data_type='s'），
以 '=+-@' 开头也不当作公式执行。有阻断冲突的导出标题明确为诊断，非正式排班。
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from models import (
    CourseEntry,
    InputSnapshot,
    Issue,
    OverrideAction,
    PersonalOverride,
    ScheduleResult,
    Task,
    Volunteer,
    make_task_id,
)
from normalize import parse_period, parse_weekday, parse_weeks


SHEET_VOLUNTEERS = "志愿者"
SHEET_COURSES = "班级课表"
SHEET_OVERRIDES = "个人修正"
SHEET_TASKS = "查课任务"

ALL_SHEETS = (SHEET_VOLUNTEERS, SHEET_COURSES, SHEET_OVERRIDES, SHEET_TASKS)


def _canon_header(value: object) -> str:
    if value is None:
        return ""
    s = str(value).replace("　", "").replace(" ", "").replace("\t", "")
    return s.lower()


def _build_alias_map(aliases: dict[str, list[str]]) -> dict[str, str]:
    m: dict[str, str] = {}
    for field, names in aliases.items():
        for n in names:
            m[_canon_header(n)] = field
    return m


_VOL_ALIASES = _build_alias_map(
    {
        "student_no": ["student_no", "studentno", "学号", "志愿者学号"],
        "name": ["name", "姓名", "志愿者姓名"],
        "class_id": ["class_id", "classid", "班级", "班级编号", "所在班级", "志愿者班级"],
    }
)
_COURSE_ALIASES = _build_alias_map(
    {
        "course_id": ["course_id", "courseid", "课程编号", "课程号", "课程id"],
        "class_id": ["class_id", "classid", "班级", "班级编号", "上课班级"],
        "course_name": ["course_name", "coursename", "课程名称", "课程名", "课程"],
        "location": ["location", "上课地点", "地点", "教室"],
        "weeks": ["weeks", "周次", "上课周次", "周"],
        "weekday": ["weekday", "星期", "周几", "星期几", "上课星期"],
        "period": ["period", "节次", "大节", "上课节次", "小节", "节"],
    }
)
_OVERRIDE_ALIASES = _build_alias_map(
    {
        "override_id": ["override_id", "overrideid", "修正编号", "编号", "修正号"],
        "student_no": ["student_no", "studentno", "学号", "志愿者学号"],
        "action": ["action", "类型", "修正类型", "动作", "操作"],
        "course_id": ["course_id", "courseid", "课程编号", "课程号"],
        "weeks": ["weeks", "周次"],
        "weekday": ["weekday", "星期", "周几", "星期几"],
        "period": ["period", "节次", "大节"],
        "remark": ["remark", "备注", "说明"],
    }
)
_TASK_ALIASES = _build_alias_map(
    {
        "task_id": ["task_id", "taskid", "任务编号", "任务号"],
        "target_class_id": [
            "target_class_id",
            "targetclassid",
            "目标班级",
            "被查班级",
            "班级",
            "检查班级",
        ],
        "weeks": ["weeks", "周次"],
        "weekday": ["weekday", "星期", "周几", "星期几"],
        "period": ["period", "节次", "大节"],
        "location": ["location", "地点", "教室", "位置"],
        "required_people": [
            "required_people",
            "requiredpeople",
            "需要人数",
            "人数",
            "需求人数",
        ],
    }
)

_VOL_REQUIRED = ("student_no", "name", "class_id")
_COURSE_REQUIRED = ("course_id", "class_id", "course_name", "weeks", "weekday", "period")
_OVERRIDE_REQUIRED = ("override_id", "student_no", "action", "weeks", "weekday", "period")
_TASK_REQUIRED = ("task_id", "target_class_id", "weeks", "weekday", "period", "required_people")


# ---------------- 模板 ----------------

_TEMPLATE_HEADERS: dict[str, list[str]] = {
    SHEET_VOLUNTEERS: ["学号", "姓名", "班级"],
    SHEET_COURSES: ["课程编号", "班级", "课程名称", "周次", "星期", "节次", "上课地点"],
    SHEET_OVERRIDES: ["修正编号", "学号", "类型", "课程编号", "周次", "星期", "节次", "备注"],
    SHEET_TASKS: ["任务编号", "目标班级", "周次", "星期", "节次", "地点", "需要人数"],
}

# 模板示例行（total_weeks=16 下合法，便于用户对照格式）。
_TEMPLATE_EXAMPLES: dict[str, list[list[object]]] = {
    SHEET_VOLUNTEERS: [["V001", "张三", "A"], ["V002", "李四", "B"]],
    SHEET_COURSES: [
        ["C001", "A", "高等数学", "1-8", "周一", "1-2"],
        ["C002", "A", "大学英语", "1-8", "周一", "1-2"],
    ],
    SHEET_OVERRIDES: [
        ["O001", "V001", "CANCEL_COURSE", "C001", "1-8", "周一", "1-2", "示例：取消其中一门"],
    ],
    SHEET_TASKS: [
        ["T001", "B", "3", "周一", "1-2", "教学楼101", 2],
    ],
}

# 模板中按文本格式存储的列（保留前导零），以 1-based 列索引表示。
_TEMPLATE_TEXT_COLS: dict[str, set[int]] = {
    SHEET_VOLUNTEERS: {1},
    SHEET_COURSES: {1},
    SHEET_OVERRIDES: {1, 2, 4},
    SHEET_TASKS: {1},
}


def create_template_workbook() -> bytes:
    """生成四表标准模板（含表头与示例行），返回 xlsx 字节。"""
    wb = Workbook()
    # openpyxl 默认有一张 Sheet，复用为第一张表。
    first = True
    for sheet in ALL_SHEETS:
        ws = wb.active if first else wb.create_sheet(title=sheet)
        first = False
        ws.title = sheet
        headers = _TEMPLATE_HEADERS[sheet]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in _TEMPLATE_EXAMPLES.get(sheet, []):
            ws.append(list(row))
        text_cols = _TEMPLATE_TEXT_COLS.get(sheet, set())
        for r in range(2, ws.max_row + 1):
            for c in text_cols:
                cell = ws.cell(row=r, column=c)
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
        widths = {"A": 14, "B": 14, "C": 16, "D": 14, "E": 12, "F": 12, "G": 14, "H": 20}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- 读取 ----------------

def _to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


def _parse_action(value: object) -> OverrideAction | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    s = str(value).strip().upper().replace("　", "").replace(" ", "")
    if s in ("ADD", "新增", "增加", "增加占用", "新增占用"):
        return OverrideAction.ADD
    if s in ("CANCEL_COURSE", "CANCEL", "REMOVE", "取消", "取消课程", "取消占用"):
        return OverrideAction.CANCEL_COURSE
    # 兼容小写与常见写法
    low = str(value).strip().lower().replace(" ", "")
    if low in ("add",):
        return OverrideAction.ADD
    if low in ("cancel_course", "cancel", "remove"):
        return OverrideAction.CANCEL_COURSE
    return None


def _parse_required_people(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 1 else None
    if isinstance(value, float):
        if value.is_integer() and int(value) >= 1:
            return int(value)
        return None
    s = str(value).strip().replace("　", "").replace(" ", "")
    if not s:
        return None
    # 拒绝小数/公式残留
    if s.startswith("="):
        return None
    if s.isdigit():
        v = int(s)
        return v if v >= 1 else None
    return None


def read_workbook(
    data: bytes, total_weeks: int, term_id: str = "2026-autumn",
    slot_mapping: dict[int, int] | None = None,
) -> tuple[InputSnapshot | None, list[Issue]]:
    """读取标准工作簿并展开为领域对象。

    成功返回 ``(snapshot, warnings)``；失败返回 ``(None, blocking_errors)``。
    ``slot_mapping`` 为空则用默认 1..10 小节映射；学校晚上 11-12 节时传入
    ``school_grid.SCHOOL_SLOT_MAPPING``，并须同步记入 ``Rules.slot_mapping``。
    """
    issues: list[Issue] = []
    if not isinstance(total_weeks, int) or total_weeks < 1:
        raise ValueError(f"非法学期总周数：{total_weeks!r}")
    if not isinstance(data, (bytes, bytearray)) or not data:
        return None, [Issue(code="INVALID_WORKBOOK", message="文件为空或格式错误")]

    try:
        wb = load_workbook(filename=io.BytesIO(bytes(data)), data_only=False)
    except Exception as exc:
        return None, [
            Issue(code="INVALID_WORKBOOK", message=f"无法解析 xlsx：{exc}")
        ]

    for sheet in ALL_SHEETS:
        if sheet not in wb.sheetnames:
            issues.append(
                Issue(code="MISSING_SHEET", message=f"缺少工作表：{sheet}", sheet=sheet)
            )
    if issues:
        return None, issues

    # 通用表读取：返回 {field: col_idx} 与行数据。
    parsed_vol_rows: list[dict] = []
    parsed_course_rows: list[dict] = []
    parsed_override_rows: list[dict] = []
    parsed_task_rows: list[dict] = []

    ok = True
    ok &= _read_sheet(
        wb[SHEET_VOLUNTEERS], SHEET_VOLUNTEERS, _VOL_ALIASES, _VOL_REQUIRED,
        issues, parsed_vol_rows, total_weeks, slot_mapping,
    )
    ok &= _read_sheet(
        wb[SHEET_COURSES], SHEET_COURSES, _COURSE_ALIASES, _COURSE_REQUIRED,
        issues, parsed_course_rows, total_weeks, slot_mapping,
    )
    ok &= _read_sheet(
        wb[SHEET_OVERRIDES], SHEET_OVERRIDES, _OVERRIDE_ALIASES, _OVERRIDE_REQUIRED,
        issues, parsed_override_rows, total_weeks, slot_mapping,
    )
    ok &= _read_sheet(
        wb[SHEET_TASKS], SHEET_TASKS, _TASK_ALIASES, _TASK_REQUIRED,
        issues, parsed_task_rows, total_weeks, slot_mapping,
    )
    if not ok:
        return None, issues

    # 跨表引用检查。
    volunteer_ids = {r["student_no"] for r in parsed_vol_rows}
    course_ids = {r["course_id"] for r in parsed_course_rows}
    course_classes: dict[str, set[str]] = {}
    for r in parsed_course_rows:
        course_classes.setdefault(r["course_id"], set()).add(r["class_id"])
    vol_class: dict[str, str] = {r["student_no"]: r["class_id"] for r in parsed_vol_rows}
    known_classes = {r["class_id"] for r in parsed_vol_rows} | {
        r["class_id"] for r in parsed_course_rows
    }

    cross_ok = True
    for r in parsed_override_rows:
        row_no: int = r["_row"]
        if r["student_no"] not in volunteer_ids:
            issues.append(
                Issue(
                    code="UNKNOWN_VOLUNTEER",
                    message=f"个人修正引用未知学号：{r['student_no']}",
                    sheet=SHEET_OVERRIDES, row=row_no, field="student_no",
                    volunteer_id=r["student_no"],
                )
            )
            cross_ok = False
        if r["action"] == OverrideAction.CANCEL_COURSE:
            cid = r["course_id"]
            if cid not in course_ids:
                issues.append(
                    Issue(
                        code="UNKNOWN_COURSE",
                        message=f"取消引用不存在的课程：{cid}",
                        sheet=SHEET_OVERRIDES, row=row_no, field="course_id",
                        volunteer_id=r["student_no"],
                    )
                )
                cross_ok = False
            elif r["student_no"] in vol_class:
                own = vol_class[r["student_no"]]
                if own not in course_classes.get(cid, set()):
                    issues.append(
                        Issue(
                            code="CANCEL_NOT_OWN_COURSE",
                            message=f"只能取消本人所在班级课程：{cid} 不属于 {own}",
                            sheet=SHEET_OVERRIDES, row=row_no, field="course_id",
                            volunteer_id=r["student_no"],
                        )
                    )
                    cross_ok = False
    for r in parsed_task_rows:
        if r["target_class_id"] not in known_classes:
            issues.append(
                Issue(
                    code="UNKNOWN_CLASS",
                    message=f"任务目标班级不存在：{r['target_class_id']}",
                    sheet=SHEET_TASKS, row=r["_row"], field="target_class_id",
                    task_id=r["task_id"],
                )
            )
            cross_ok = False
    if not cross_ok:
        return None, issues

    # 展开为领域对象（确定性排序，与 Excel 行顺序无关）。
    warnings: list[Issue] = []

    volunteers = tuple(
        sorted(
            (
                Volunteer(
                    id=r["student_no"],
                    student_no=r["student_no"],
                    name=r["name"],
                    class_id=r["class_id"],
                )
                for r in parsed_vol_rows
            ),
            key=lambda v: v.student_no,
        )
    )

    courses: list[CourseEntry] = []
    for r in sorted(parsed_course_rows, key=lambda x: (x["course_id"], x["_row"])):
        for w in r["weeks_list"]:
            for p in r["period_list"]:
                courses.append(
                    CourseEntry(
                        id=r["course_id"],
                        class_id=r["class_id"],
                        slot=(w, r["weekday"], p),
                        course_name=r["course_name"],
                        location=r.get('location', ''),
                    )
                )
    courses_sorted = tuple(sorted(courses, key=lambda c: (c.id, c.slot)))

    overrides: list[PersonalOverride] = []
    for r in sorted(parsed_override_rows, key=lambda x: (x["override_id"], x["_row"])):
        for w in r["weeks_list"]:
            for p in r["period_list"]:
                overrides.append(
                    PersonalOverride(
                        id=r["override_id"],
                        volunteer_id=r["student_no"],
                        action=r["action"],
                        course_id=r["course_id"] if r["action"] == OverrideAction.CANCEL_COURSE else None,
                        slot=(w, r["weekday"], p),
                    )
                )
    overrides_sorted = tuple(sorted(overrides, key=lambda o: (o.id, o.slot)))

    tasks: list[Task] = []
    for r in sorted(parsed_task_rows, key=lambda x: (x["task_id"], x["_row"])):
        for w in r["weeks_list"]:
            slot = (w, r["weekday"], r["period_list"][0])
            tasks.append(
                Task(
                    id=make_task_id(r["task_id"], slot),
                    target_class_id=r["target_class_id"],
                    slot=slot,
                    location=r.get("location") or "",
                    required_people=r["required_people"],
                    raw_id=r["task_id"],
                )
            )
    tasks_sorted = tuple(sorted(tasks, key=lambda t: t.id))

    # 重复占用提示（不同编号、完全相同的 ADD 占用）：非阻断。
    seen_add: dict[tuple[str, tuple], str] = {}
    for o in overrides_sorted:
        if o.action != OverrideAction.ADD:
            continue
        key = (o.volunteer_id, o.slot)
        if key in seen_add and seen_add[key] != o.id:
            warnings.append(
                Issue(
                    code="DUPLICATE_OCCUPANCY",
                    message=f"重复占用：{o.volunteer_id} 在 {o.slot} 已有 {seen_add[key]}，{o.id} 时间不重复计数",
                    sheet=SHEET_OVERRIDES, field="period",
                    volunteer_id=o.volunteer_id,
                )
            )
        else:
            seen_add[key] = o.id

    snapshot = InputSnapshot(
        term_id=term_id,
        total_weeks=total_weeks,
        volunteers=volunteers,
        courses=courses_sorted,
        overrides=overrides_sorted,
        tasks=tasks_sorted,
    )
    return snapshot, warnings


def _read_sheet(
    ws,
    sheet_name: str,
    alias_map: dict[str, str],
    required: tuple[str, ...],
    issues: list[Issue],
    out_rows: list[dict],
    total_weeks: int,
    slot_mapping: dict[int, int] | None = None,
) -> bool:
    """读取单表。成功返回 True；失败追加 Issue 并返回 False。"""
    # 表头（第 1 行）。
    header_cells = list(ws[1])
    col_to_field: dict[int, str] = {}
    seen_fields: dict[str, int] = {}
    header_ok = True
    for idx, cell in enumerate(header_cells, start=1):
        if cell.data_type == "f":
            issues.append(
                Issue(code="FORMULA_CELL", message="表头不能是公式",
                      sheet=sheet_name, row=1, field=str(cell.value))
            )
            header_ok = False
            continue
        canon = _canon_header(cell.value)
        if not canon:
            continue
        field = alias_map.get(canon)
        if field is None:
            continue  # 未知列忽略，允许备注列等扩展
        if field in seen_fields:
            issues.append(
                Issue(code="DUPLICATE_COLUMN", message=f"重复列：{cell.value}",
                      sheet=sheet_name, row=1, field=field)
            )
            header_ok = False
        else:
            seen_fields[field] = idx
            col_to_field[idx] = field
    for f in required:
        if f not in seen_fields:
            issues.append(
                Issue(code="MISSING_COLUMN", message=f"缺少列：{f}",
                      sheet=sheet_name, row=1, field=f)
            )
            header_ok = False
    if not header_ok:
        return False

    field_to_col = {f: c for c, f in col_to_field.items()}
    # 唯一性跟踪。
    seen_vol: set[str] = set()
    seen_override: set[str] = set()
    seen_task: set[str] = set()

    sheet_ok = True
    for row_no in range(2, ws.max_row + 1):
        # 整行空则跳过。
        all_empty = True
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_no, column=col_idx)
            v = cell.value
            if v is not None and (not isinstance(v, str) or v.strip() != ""):
                all_empty = False
                break
        if all_empty:
            continue

        # 公式检查：行内任一已映射列为公式即阻断。
        formula_field = None
        for f, c in field_to_col.items():
            cell = ws.cell(row=row_no, column=c)
            if cell.data_type == "f":
                formula_field = f
                break
        if formula_field is not None:
            issues.append(
                Issue(code="FORMULA_CELL", message="不接受公式输入，请粘贴为值",
                      sheet=sheet_name, row=row_no, field=formula_field)
            )
            sheet_ok = False
            continue

        raw: dict[str, object] = {}
        for f, c in field_to_col.items():
            raw[f] = ws.cell(row=row_no, column=c).value

        if sheet_name == SHEET_VOLUNTEERS:
            if not _parse_vol_row(raw, row_no, issues, out_rows, seen_vol):
                sheet_ok = False
        elif sheet_name == SHEET_COURSES:
            if not _parse_course_row(raw, row_no, issues, out_rows, total_weeks, slot_mapping):
                sheet_ok = False
        elif sheet_name == SHEET_OVERRIDES:
            if not _parse_override_row(raw, row_no, issues, out_rows, seen_override, total_weeks, slot_mapping):
                sheet_ok = False
        elif sheet_name == SHEET_TASKS:
            if not _parse_task_row(raw, row_no, issues, out_rows, seen_task, total_weeks, slot_mapping):
                sheet_ok = False
    return sheet_ok


def _parse_vol_row(raw, row_no, issues, out_rows, seen) -> bool:
    student_no = _to_text(raw.get("student_no"))
    name = _to_text(raw.get("name"))
    class_id = _to_text(raw.get("class_id"))
    ok = True
    if not student_no:
        issues.append(Issue(code="EMPTY_FIELD", message="学号不能为空",
                            sheet=SHEET_VOLUNTEERS, row=row_no, field="student_no"))
        ok = False
    if not name:
        issues.append(Issue(code="EMPTY_FIELD", message="姓名不能为空",
                            sheet=SHEET_VOLUNTEERS, row=row_no, field="name",
                            volunteer_id=student_no or None))
        ok = False
    if not class_id:
        issues.append(Issue(code="EMPTY_FIELD", message="班级不能为空",
                            sheet=SHEET_VOLUNTEERS, row=row_no, field="class_id",
                            volunteer_id=student_no or None))
        ok = False
    if not ok:
        return False
    if student_no in seen:
        issues.append(Issue(code="DUPLICATE_ID", message=f"学号重复：{student_no}",
                            sheet=SHEET_VOLUNTEERS, row=row_no, field="student_no",
                            volunteer_id=student_no))
        return False
    seen.add(student_no)
    out_rows.append({"student_no": student_no, "name": name, "class_id": class_id, "_row": row_no})
    return True


def _parse_course_row(raw, row_no, issues, out_rows, total_weeks, slot_mapping=None) -> bool:
    course_id = _to_text(raw.get("course_id"))
    class_id = _to_text(raw.get("class_id"))
    course_name = _to_text(raw.get("course_name"))
    weeks_raw = raw.get("weeks")
    weekday_raw = raw.get("weekday")
    period_raw = raw.get("period")
    ok = True
    if not course_id:
        issues.append(Issue(code="EMPTY_FIELD", message="课程编号不能为空",
                            sheet=SHEET_COURSES, row=row_no, field="course_id"))
        ok = False
    if not class_id:
        issues.append(Issue(code="EMPTY_FIELD", message="班级不能为空",
                            sheet=SHEET_COURSES, row=row_no, field="class_id"))
        ok = False
    if not course_name:
        issues.append(Issue(code="EMPTY_FIELD", message="课程名称不能为空",
                            sheet=SHEET_COURSES, row=row_no, field="course_name"))
        ok = False
    weeks_list = None
    weekday = None
    period_list = None
    if weeks_raw is None or (isinstance(weeks_raw, str) and not weeks_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="周次不能为空",
                            sheet=SHEET_COURSES, row=row_no, field="weeks"))
        ok = False
    else:
        try:
            weeks_list = parse_weeks(_to_text(weeks_raw) if not isinstance(weeks_raw, int) else weeks_raw, total_weeks)
        except ValueError as exc:
            issues.append(Issue(code="INVALID_WEEKS", message=str(exc),
                                sheet=SHEET_COURSES, row=row_no, field="weeks"))
            ok = False
    try:
        weekday = parse_weekday(_to_text(weekday_raw) if isinstance(weekday_raw, str) else weekday_raw)  # type: ignore[arg-type]
    except ValueError as exc:
        issues.append(Issue(code="INVALID_WEEKDAY", message=str(exc),
                            sheet=SHEET_COURSES, row=row_no, field="weekday"))
        ok = False
    if period_raw is None or (isinstance(period_raw, str) and not period_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="节次不能为空",
                            sheet=SHEET_COURSES, row=row_no, field="period"))
        ok = False
    else:
        try:
            period_list = parse_period(_to_text(period_raw) if isinstance(period_raw, str) else period_raw, slot_mapping)  # type: ignore[arg-type]
        except ValueError as exc:
            issues.append(Issue(code="INVALID_PERIOD", message=str(exc),
                                sheet=SHEET_COURSES, row=row_no, field="period"))
            ok = False
    if not ok:
        return False
    assert weeks_list is not None and weekday is not None and period_list is not None
    out_rows.append({
        "course_id": course_id, "class_id": class_id, "course_name": course_name,
        "weeks_list": weeks_list, "weekday": weekday, "period_list": period_list,
        "location": _to_text(raw.get('location')),
        "_row": row_no,
    })
    return True


def _parse_override_row(raw, row_no, issues, out_rows, seen, total_weeks, slot_mapping=None) -> bool:
    override_id = _to_text(raw.get("override_id"))
    student_no = _to_text(raw.get("student_no"))
    action_raw = raw.get("action")
    course_id = _to_text(raw.get("course_id"))
    weeks_raw = raw.get("weeks")
    weekday_raw = raw.get("weekday")
    period_raw = raw.get("period")
    remark = _to_text(raw.get("remark"))
    ok = True
    if not override_id:
        issues.append(Issue(code="EMPTY_FIELD", message="修正编号不能为空",
                            sheet=SHEET_OVERRIDES, row=row_no, field="override_id"))
        ok = False
    elif override_id in seen:
        issues.append(Issue(code="DUPLICATE_ID", message=f"修正编号重复：{override_id}",
                            sheet=SHEET_OVERRIDES, row=row_no, field="override_id",
                            volunteer_id=student_no or None))
        return False
    if not student_no:
        issues.append(Issue(code="EMPTY_FIELD", message="学号不能为空",
                            sheet=SHEET_OVERRIDES, row=row_no, field="student_no"))
        ok = False
    action = _parse_action(action_raw)
    if action is None:
        issues.append(Issue(code="INVALID_ACTION", message=f"类型必须为 ADD 或 CANCEL_COURSE：{action_raw!r}",
                            sheet=SHEET_OVERRIDES, row=row_no, field="action",
                            volunteer_id=student_no or None))
        ok = False
    weeks_list = None
    weekday = None
    period_list = None
    if weeks_raw is None or (isinstance(weeks_raw, str) and not weeks_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="周次不能为空",
                            sheet=SHEET_OVERRIDES, row=row_no, field="weeks",
                            volunteer_id=student_no or None))
        ok = False
    else:
        try:
            weeks_list = parse_weeks(_to_text(weeks_raw) if not isinstance(weeks_raw, int) else weeks_raw, total_weeks)
        except ValueError as exc:
            issues.append(Issue(code="INVALID_WEEKS", message=str(exc),
                                sheet=SHEET_OVERRIDES, row=row_no, field="weeks",
                                volunteer_id=student_no or None))
            ok = False
    try:
        weekday = parse_weekday(_to_text(weekday_raw) if isinstance(weekday_raw, str) else weekday_raw)  # type: ignore[arg-type]
    except ValueError as exc:
        issues.append(Issue(code="INVALID_WEEKDAY", message=str(exc),
                            sheet=SHEET_OVERRIDES, row=row_no, field="weekday",
                            volunteer_id=student_no or None))
        ok = False
    if period_raw is None or (isinstance(period_raw, str) and not period_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="节次不能为空",
                            sheet=SHEET_OVERRIDES, row=row_no, field="period",
                            volunteer_id=student_no or None))
        ok = False
    else:
        try:
            period_list = parse_period(_to_text(period_raw) if isinstance(period_raw, str) else period_raw, slot_mapping)  # type: ignore[arg-type]
        except ValueError as exc:
            issues.append(Issue(code="INVALID_PERIOD", message=str(exc),
                                sheet=SHEET_OVERRIDES, row=row_no, field="period",
                                volunteer_id=student_no or None))
            ok = False
    if action == OverrideAction.ADD and course_id:
        issues.append(Issue(code="ADD_WITH_COURSE", message="ADD 不应填写课程编号",
                            sheet=SHEET_OVERRIDES, row=row_no, field="course_id",
                            volunteer_id=student_no or None))
        ok = False
    if action == OverrideAction.CANCEL_COURSE and not course_id:
        issues.append(Issue(code="EMPTY_FIELD", message="CANCEL_COURSE 必须填写课程编号",
                            sheet=SHEET_OVERRIDES, row=row_no, field="course_id",
                            volunteer_id=student_no or None))
        ok = False
    if not ok:
        return False
    assert weeks_list is not None and weekday is not None and period_list is not None and action is not None
    seen.add(override_id)
    out_rows.append({
        "override_id": override_id, "student_no": student_no, "action": action,
        "course_id": course_id, "weeks_list": weeks_list, "weekday": weekday,
        "period_list": period_list, "remark": remark, "_row": row_no,
    })
    return True


def _parse_task_row(raw, row_no, issues, out_rows, seen, total_weeks, slot_mapping=None) -> bool:
    task_id = _to_text(raw.get("task_id"))
    target_class_id = _to_text(raw.get("target_class_id"))
    weeks_raw = raw.get("weeks")
    weekday_raw = raw.get("weekday")
    period_raw = raw.get("period")
    location = _to_text(raw.get("location"))
    req_raw = raw.get("required_people")
    ok = True
    if not task_id:
        issues.append(Issue(code="EMPTY_FIELD", message="任务编号不能为空",
                            sheet=SHEET_TASKS, row=row_no, field="task_id"))
        ok = False
    elif task_id in seen:
        issues.append(Issue(code="DUPLICATE_ID", message=f"任务编号重复：{task_id}",
                            sheet=SHEET_TASKS, row=row_no, field="task_id", task_id=task_id))
        return False
    if not target_class_id:
        issues.append(Issue(code="EMPTY_FIELD", message="目标班级不能为空",
                            sheet=SHEET_TASKS, row=row_no, field="target_class_id",
                            task_id=task_id or None))
        ok = False
    weeks_list = None
    weekday = None
    period_list = None
    if weeks_raw is None or (isinstance(weeks_raw, str) and not weeks_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="周次不能为空",
                            sheet=SHEET_TASKS, row=row_no, field="weeks",
                            task_id=task_id or None))
        ok = False
    else:
        try:
            weeks_list = parse_weeks(_to_text(weeks_raw) if not isinstance(weeks_raw, int) else weeks_raw, total_weeks)
        except ValueError as exc:
            issues.append(Issue(code="INVALID_WEEKS", message=str(exc),
                                sheet=SHEET_TASKS, row=row_no, field="weeks",
                                task_id=task_id or None))
            ok = False
    try:
        weekday = parse_weekday(_to_text(weekday_raw) if isinstance(weekday_raw, str) else weekday_raw)  # type: ignore[arg-type]
    except ValueError as exc:
        issues.append(Issue(code="INVALID_WEEKDAY", message=str(exc),
                            sheet=SHEET_TASKS, row=row_no, field="weekday",
                            task_id=task_id or None))
        ok = False
    if period_raw is None or (isinstance(period_raw, str) and not period_raw.strip()):
        issues.append(Issue(code="EMPTY_FIELD", message="节次不能为空",
                            sheet=SHEET_TASKS, row=row_no, field="period",
                            task_id=task_id or None))
        ok = False
    else:
        try:
            period_list = parse_period(_to_text(period_raw) if isinstance(period_raw, str) else period_raw, slot_mapping)  # type: ignore[arg-type]
        except ValueError as exc:
            issues.append(Issue(code="INVALID_PERIOD", message=str(exc),
                                sheet=SHEET_TASKS, row=row_no, field="period",
                                task_id=task_id or None))
            ok = False
    if period_list is not None and len(period_list) != 1:
        issues.append(Issue(code="TASK_MULTIPLE_SLOTS", message=f"查课任务只占一个标准大节：{period_raw!r}",
                            sheet=SHEET_TASKS, row=row_no, field="period",
                            task_id=task_id or None))
        ok = False
    req = _parse_required_people(req_raw)
    if req is None:
        if req_raw is None or (isinstance(req_raw, str) and not req_raw.strip()):
            issues.append(Issue(code="EMPTY_FIELD", message="需要人数不能为空",
                                sheet=SHEET_TASKS, row=row_no, field="required_people",
                                task_id=task_id or None))
        else:
            issues.append(Issue(code="INVALID_REQUIRED_PEOPLE", message=f"需要人数必须为正整数：{req_raw!r}",
                                sheet=SHEET_TASKS, row=row_no, field="required_people",
                                task_id=task_id or None))
        ok = False
    if not ok:
        return False
    assert weeks_list is not None and weekday is not None and period_list is not None and req is not None
    seen.add(task_id)
    out_rows.append({
        "task_id": task_id, "target_class_id": target_class_id,
        "weeks_list": weeks_list, "weekday": weekday, "period_list": period_list,
        "location": location, "required_people": req, "_row": row_no,
    })
    return True

# ---------------- 结果导出（M2b） ----------------

EXPORT_SHEETS = ("总排班", "个人排班", "未分配任务", "冲突报告", "次数统计")


def _set_text(ws, row: int, col: int, value: object) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"


def _set_num(ws, row: int, col: int, value: int) -> None:
    cell = ws.cell(row=row, column=col)
    cell.value = int(value)
    cell.data_type = "n"


def export_workbook(snapshot: InputSnapshot, result: ScheduleResult, *, historical: bool = False,
                    slot_mapping: dict[int, int] | None = None) -> bytes:
    """Person-first task lists; technical slot numbers never leak into user tables."""
    from collections import Counter
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from schedule_display import assignment_rows, task_details, course_index

    status = '诊断-存在阻断冲突' if result.validation.blocking_errors else '正式'
    if historical:
        status = '历史-输入已更新，请勿作为当前排班'
    people = assignment_rows(snapshot, result.assignments, slot_mapping)
    columns = ['志愿者','学号','志愿者班级','周次','星期','节次','检查班级','课程','上课地点','任务说明']
    tasks = {t.id: t for t in snapshot.tasks}
    index = course_index(snapshot)
    shortages = []
    for shortage in sorted(result.shortages, key=lambda s: (tasks[s.task_id].slot if s.task_id in tasks else (0,0,0), s.task_id)):
        task = tasks.get(shortage.task_id)
        if task:
            shortages.append({**task_details(snapshot, task, slot_mapping, index),
                              '需要人数':shortage.required, '实排人数':shortage.assigned,
                              '缺少人数':shortage.missing, '说明':'本次未找到足够可用人员'})
    issues = [{'类型':kind, '编码':i.code, '说明':i.message, '任务':i.task_id or '',
               '志愿者':i.volunteer_id or ''}
              for kind, entries in [('阻断',result.validation.blocking_errors),('警告',result.validation.warnings)]
              for i in entries]
    counts = Counter(a.volunteer_id for a in result.assignments)
    stats = [{'学号':v.student_no,'姓名':v.name,'志愿者班级':v.class_id,'分配次数':counts[v.id]}
             for v in sorted(snapshot.volunteers,key=lambda v:(v.name,v.student_no))]
    sheets = [('总排班',columns,people), ('个人排班',columns,people),
              ('未分配任务',['周次','星期','节次','检查班级','课程','上课地点','需要人数','实排人数','缺少人数','说明'],shortages),
              ('冲突报告',['类型','编码','说明','任务','志愿者'],issues),
              ('次数统计',['学号','姓名','志愿者班级','分配次数'],stats)]
    wb = Workbook()
    widths = {'志愿者':16,'学号':20,'志愿者班级':20,'周次':12,'星期':12,'节次':14,
              '检查班级':20,'课程':32,'上课地点':28,'任务说明':85,'说明':60,'编码':30,'任务':30}
    for number,(name,headers,rows) in enumerate(sheets):
        ws = wb.active if number == 0 else wb.create_sheet(name)
        ws.title = name
        ws.cell(1,1,f'{name}（{status}）').font = Font(bold=True,size=14)
        for col,header in enumerate(headers,1):
            cell = ws.cell(2,col,header)
            cell.font = Font(bold=True,color='FFFFFF')
            cell.fill = PatternFill('solid',fgColor='294E70')
            ws.column_dimensions[get_column_letter(col)].width = widths.get(header,14)
        for r,values in enumerate(rows,3):
            for col,header in enumerate(headers,1):
                value = values.get(header,'')
                if isinstance(value,int):
                    _set_num(ws,r,col,value)
                else:
                    _set_text(ws,r,col,value)
                ws.cell(r,col).alignment = Alignment(vertical='center',wrap_text=True)
                if r % 2 == 0:
                    ws.cell(r,col).fill = PatternFill('solid',fgColor='F0F4F8')
        ws.freeze_panes = 'D3' if name in ('总排班','个人排班') else 'A3'
        ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{max(2,ws.max_row)}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
