"""师生绑定码共用表格样式；全部按文本写入，保护前导零并避免公式解释。"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def build_binding_workbook(
    *,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    code_column: int,
    widths: list[int],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    sheet.append(headers)
    border = Border(
        **{side: Side(style="thin", color="D9D9D9") for side in ("left", "right", "top", "bottom")}
    )
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1890FF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28
    for values in rows:
        sheet.append(values)
        sheet.row_dimensions[sheet.max_row].height = 22
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.border = border
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="center")
        sheet.cell(sheet.max_row, code_column).font = Font(
            name="Consolas", size=12, bold=True, color="D4380D"
        )
        sheet.cell(sheet.max_row, code_column).alignment = Alignment(
            horizontal="center", vertical="center"
        )
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
