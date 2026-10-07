"""身份模块服务层：授权、业务规则与事务边界（技术方案 5.1、7）。

- 使用同步 SQLAlchemy Session；由 Router 依赖注入的 `get_db` 提供，
  服务在成功路径显式提交，失败抛 AppError 由依赖回滚关闭。
- 访问令牌携带会话引用；每次请求回查 auth_session 有效性。
- 刷新凭证轮换：旧凭证即时失效，检测重放则撤销整个刷新族。
- 绑定码核销与学生绑定在同一原子事务完成，唯一约束阻止重复绑定。
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from datetime import datetime, timedelta

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.database import utcnow
from app.core.exceptions import (
    AppError,
    ConflictError,
    ErrorCode,
    NotFoundError,
    PermissionDeniedError,
    UnauthenticatedError,
)
from app.core.permissions import PermissionCode, RoleCode
from app.core.security import (
    create_access_token,
    generate_friendly_binding_code,
    generate_secure_token,
    hash_token,
    verify_password,
)
from app.modules.academic.models import Student
from app.modules.audit.models import AuditLog
from app.modules.identity.models import (
    AuthSession,
    BindingTokenStatus,
    ClientType,
    UserAccount,
    UserStatus,
    WechatIdentity,
)
from app.modules.identity.policy import (
    OPTIONAL_PERMISSION_CODES,
    assignable_roles,
    can_change_roles,
)
from app.modules.identity.repository import IdentityRepository
from app.modules.identity.schemas import (
    OptionalPermissionState,
    OptionalPermissionTargetItem,
    RoleAssignmentTargetItem,
    TokenPair,
)
from app.modules.identity.wechat import (
    WechatCode2SessionClient,
    build_default_client,
)


@dataclass(frozen=True)
class CurrentUser:
    """服务端会话解析出的当前身份，供权限依赖使用。

    pre_binding=True 表示"微信登录成功但尚未绑定 student_id"的受限会话状态
    （PERMISSIONS.md 1.6）：不是角色，仅允许身份白名单接口。Web 管理员即使
    未绑学生也因持有角色而不进入该状态。
    """

    id: int
    username: str | None
    display_name: str
    status: str
    roles: list[str]
    permissions: list[str]
    pre_binding: bool = False

    def has_role(self, role: str) -> bool:
        return role in self.roles

    def has_permission(self, code: str) -> bool:
        return code in self.permissions


class IdentityService:
    def __init__(
        self,
        session: Session,
        *,
        wechat_client: WechatCode2SessionClient | None = None,
    ) -> None:
        self._session = session
        self._repo = IdentityRepository(session)
        self._settings = get_settings()
        self._wechat_client = wechat_client

    # ------------------------------------------------------------------ #
    # 令牌与会话
    # ------------------------------------------------------------------ #
    def _issue_session(
        self, user: UserAccount, client_type: ClientType
    ) -> tuple[AuthSession, str, str]:
        refresh = generate_secure_token(32)
        family = str(uuid.uuid4())
        session = AuthSession(
            user_id=user.id,
            client_type=client_type.value,
            refresh_token_hash=hash_token(refresh),
            refresh_family_id=family,
            expires_at=utcnow() + timedelta(days=self._settings.session_expire_days),
        )
        self._session.add(session)
        self._session.flush()  # 需要主键：不 commit，交由外层事务

        access = create_access_token(
            user_id=user.id,
            session_id=session.id,
            roles=self._repo.list_role_codes(user.id),
            permissions=self._repo.list_effective_permissions(user.id),
        )
        return session, access, refresh

    def _token_pair(self, access: str, refresh: str) -> TokenPair:
        return TokenPair(
            access_token=access,
            expires_in=self._settings.access_token_expire_minutes * 60,
            refresh_token=refresh,
        )

    # ------------------------------------------------------------------ #
    # Web 账号密码登录
    # ------------------------------------------------------------------ #
    def login_web(self, username: str, password: str) -> TokenPair:
        user = self._repo.get_user_by_username(username)
        if user is None:
            raise UnauthenticatedError("账号或密码错误")

        now = utcnow()
        if user.locked_until is not None and user.locked_until > now:
            raise AppError(
                ErrorCode.STATE_CONFLICT,
                "账号已锁定，请稍后重试",
                http_status=429,
            )
        if not user.is_active:
            raise UnauthenticatedError("账号未启用")
        if user.password_hash is None or not verify_password(password, user.password_hash):
            self._register_login_failure(user, now)
            self._session.commit()  # 锁定计数需独立落地
            raise UnauthenticatedError("账号或密码错误")

        user.failed_login_count = 0
        user.locked_until = None
        _session, access, refresh = self._issue_session(user, ClientType.WEB)
        self._session.commit()
        return self._token_pair(access, refresh)

    def _register_login_failure(self, user: UserAccount, now) -> None:
        user.failed_login_count += 1
        if user.failed_login_count >= self._settings.login_max_failed:
            user.locked_until = now + timedelta(minutes=self._settings.login_lock_minutes)
            user.failed_login_count = 0

    # ------------------------------------------------------------------ #
    # 微信登录（P2）：code2session → appid+openid 定位/建号 → 受限会话
    # ------------------------------------------------------------------ #
    def login_wechat(self, code: str) -> tuple[TokenPair, bool]:
        """微信 code 换取身份并签发会话。

        返回 (令牌对, need_binding)。首次登录的 openid 自动建一个"无口令、无角色"
        账号并绑定微信身份；该账号因未绑定 student_id 而进入 PRE_BINDING 受限会话，
        由调用方据 need_binding 引导前端跳转绑定页。STUDENT 角色在绑定成功时才自动授予。
        code2session 的凭证无效/上游错误/超时由客户端映射为对应 AppError 直接冒泡。
        """
        if not self._settings.wechat_login_enabled:
            raise PermissionDeniedError("微信登录当前未开放")
        client = self._wechat_client or build_default_client()
        outcome = client.exchange(code)  # 失败即抛出已映射的错误
        appid = self._settings.wechat_appid

        user = self._repo.get_user_by_wechat(appid, outcome.openid)
        created = False
        if user is None:
            user = UserAccount(
                username=None,
                password_hash=None,
                display_name="微信用户",
                status=UserStatus.ACTIVE.value,
            )
            self._session.add(user)
            self._session.flush()  # 取得主键以关联微信身份
            self._repo.add_wechat_identity(user.id, appid, outcome.openid)
            created = True
        elif not user.is_active:
            raise UnauthenticatedError("账号未启用")

        # 新账号的可用性已由构造保证；既有账号已在上分支校验。
        _session, access, refresh = self._issue_session(user, ClientType.WECHAT)

        need_binding = user.student_id is None and not self._repo.list_role_codes(user.id)
        if created:
            self._record_audit(
                actor_user_id=user.id,
                action="identity.wechat_new_account",
                resource_type="user_account",
                resource_id=str(user.id),
                before=None,
                after={"source": "wechat", "need_binding": need_binding},
                reason=None,
                request_id=None,
            )
        self._session.commit()  # 建号 + 微信身份 + 会话同事务单次提交
        return self._token_pair(access, refresh), need_binding

    # ------------------------------------------------------------------ #
    # 刷新（轮换 + 重放检测）
    # ------------------------------------------------------------------ #
    def refresh(self, refresh_token: str) -> TokenPair:
        current = self._repo.get_session_by_refresh_hash(hash_token(refresh_token))
        if current is None:
            raise UnauthenticatedError("刷新凭证无效")

        # 重放检测：已撤销的凭证被再次使用 → 撤销整个刷新族。
        if current.revoked_at is not None:
            self._revoke_family(current.refresh_family_id)
            self._session.commit()
            raise UnauthenticatedError("检测到凭证异常复用，会话已失效")

        if current.expires_at <= utcnow():
            raise UnauthenticatedError("会话已过期")

        user = self._repo.get_user_by_id(current.user_id)
        if user is None or not user.is_active:
            raise UnauthenticatedError("账号不可用")

        # 轮换：撤销旧凭证，签发同族新凭证。
        current.revoked_at = utcnow()
        new_refresh = generate_secure_token(32)
        new_session = AuthSession(
            user_id=user.id,
            client_type=current.client_type,
            refresh_token_hash=hash_token(new_refresh),
            refresh_family_id=current.refresh_family_id,
            expires_at=utcnow() + timedelta(days=self._settings.session_expire_days),
        )
        self._session.add(new_session)
        self._session.flush()

        access = create_access_token(
            user_id=user.id,
            session_id=new_session.id,
            roles=self._repo.list_role_codes(user.id),
            permissions=self._repo.list_effective_permissions(user.id),
        )
        self._session.commit()
        return self._token_pair(access, new_refresh)

    def _revoke_family(self, family_id: str) -> None:
        now = utcnow()
        for s in self._repo.sessions_in_family(family_id):
            if s.revoked_at is None:
                s.revoked_at = now

    # ------------------------------------------------------------------ #
    # 登出
    # ------------------------------------------------------------------ #
    def logout(self, session_id: int) -> None:
        session = self._repo.get_session_by_id(session_id)
        if session is not None and session.revoked_at is None:
            session.revoked_at = utcnow()
            self._session.commit()

    # ------------------------------------------------------------------ #
    # 会话解析（供权限依赖调用）
    # ------------------------------------------------------------------ #
    def resolve_session_user(self, user_id: int, session_id: int) -> CurrentUser:
        user = self._repo.get_user_by_id(user_id)
        if user is None or not user.is_active:
            raise UnauthenticatedError("账号不可用")
        session = self._repo.get_session_by_id(session_id)
        if session is None or not session.is_valid:
            raise UnauthenticatedError("会话已失效，请重新登录")
        # 会话与身份必须自洽（P1 第 6 步）：令牌 sub 与 sid 指向不同账号一律拒绝，
        # 防止被篡改/伪造的 (user_id, session_id) 组合冒用他人身份。
        if session.user_id != user_id:
            raise UnauthenticatedError("会话与身份不匹配")

        session.last_used_at = utcnow()
        self._session.commit()  # 刷新 last_used，只读事务下的轻量写入
        roles = self._repo.list_role_codes(user.id)
        # PRE_BINDING：微信侧登录成功但尚未完成学生绑定（无 student_id 且无任何角色）。
        pre_binding = user.student_id is None and not roles
        return CurrentUser(
            id=user.id,
            username=user.username,
            display_name=user.display_name,
            status=user.status,
            roles=roles,
            permissions=self._repo.list_effective_permissions(user.id),
            pre_binding=pre_binding,
        )

    # ------------------------------------------------------------------ #
    # 授权管理闭环（P1 步骤 5）：同事务加锁、事务内重验、审计同事务、单次提交
    # ------------------------------------------------------------------ #
    def _lock_accounts_ordered(
        self, actor_user_id: int, target_user_id: int
    ) -> dict[int, UserAccount]:
        """按 ID 升序锁定操作者与目标账号（固定顺序防死锁），返回 id->UserAccount。

        同一事务内加锁 + populate_existing 读取，串行化对同一账号的并发变更，
        并读到已提交事务刷新过的 lock_version。
        """
        locked: dict[int, UserAccount] = {}
        for uid in sorted({actor_user_id, target_user_id}):
            user = self._repo.get_user_by_id_for_update(uid)
            if user is None:
                if uid == actor_user_id:
                    raise UnauthenticatedError("操作者账号不可用")
                raise NotFoundError("目标账号不存在")
            locked[uid] = user
        return locked

    def _require_actor_permission(self, actor_user_id: int, code: str) -> None:
        """事务内重新读取操作者有效权限并校验（P1 步骤 5 / 13.2 纵深防御）。"""
        perms = set(self._repo.list_effective_permissions(actor_user_id))
        if code not in perms:
            raise PermissionDeniedError()

    def _record_audit(
        self,
        *,
        actor_user_id: int,
        action: str,
        resource_type: str,
        resource_id: str,
        before: dict[str, object] | None,
        after: dict[str, object] | None,
        reason: str | None,
        request_id: str | None,
    ) -> None:
        """追加审计事实（同事务，仅 flush；提交由公共方法统一负责）。"""
        self._session.add(
            AuditLog(
                actor_user_id=actor_user_id,
                action=action,
                resource_type=resource_type,
                resource_id=resource_id,
                before_json=before,
                after_json=after,
                reason=reason,
                request_id=request_id,
            )
        )
        self._session.flush()

    def assign_roles(
        self,
        *,
        actor_user_id: int,
        target_user_id: int,
        desired_codes: list[str],
        expected_version: int,
        reason: str | None,
        request_id: str | None,
    ) -> tuple[list[str], int]:
        """完整替换目标角色集合，差集校验 + 乐观并发 + 同事务审计。

        错误映射：未知角色 code 422；缺 role.assign 或越边界 403；目标不可见 404；
        版本不符 409。返回实际角色集合与新 lock_version。
        """
        from app.modules.identity.account_service import lock_admin_registry, require_another_admin

        self._session.rollback()
        lock_admin_registry(self._session)
        locked = self._lock_accounts_ordered(actor_user_id, target_user_id)
        target = locked[target_user_id]
        if not locked[actor_user_id].is_active:
            raise UnauthenticatedError("操作者账号已停用")

        # 事务内重读有效身份：操作者角色、目标当前角色。
        actor_roles = frozenset(self._repo.list_role_codes(actor_user_id))
        current_roles = frozenset(self._repo.list_role_codes(target_user_id))

        # 未知 code → 422（schema 之外再兜一层，错误信息统一）。
        known = frozenset(r.value for r in RoleCode)
        desired = frozenset(desired_codes)
        if not desired <= known:
            raise AppError(
                ErrorCode.VALIDATION_ERROR,
                "包含未知角色代码",
                http_status=422,
                field_errors={"roles": sorted(desired - known)},
            )

        # 权限与边界校验（顺序：先功能权限 403，再角色差集边界 403）。
        self._require_actor_permission(actor_user_id, PermissionCode.ROLE_ASSIGN.value)
        if not can_change_roles(actor_roles, current_roles, desired):
            raise PermissionDeniedError("越权或受保护角色的变更被拒绝")

        # 乐观并发：版本不符 → 409。
        if target.lock_version != expected_version:
            raise ConflictError(
                ErrorCode.VERSION_CONFLICT, "目标已被并发修改，请刷新后重试"
            )

        before_roles = sorted(current_roles)
        if (
            RoleCode.SUPER_ADMIN.value in current_roles - desired
            and target.status == UserStatus.ACTIVE.value
        ):
            require_another_admin(self._session, target.id)
        if RoleCode.STUDENT_AFFAIRS_MANAGER.value in desired - current_roles:
            student = (
                self._session.execute(
                    select(Student).where(Student.id == target.student_id)
                    .with_for_update().execution_options(populate_existing=True)
                ).scalar_one_or_none()
                if target.student_id is not None else None
            )
            if (
                target.status != UserStatus.ACTIVE.value
                or student is None or student.status != "ACTIVE"
            ):
                raise ConflictError(
                    ErrorCode.STATE_CONFLICT, "负责人须使用已绑定有效学生的启用账号"
                )
        # 应用：把 user_role 关联精确对齐 desired（SQLAlchemy 生成最小增删）。
        role_objs = [self._repo.get_or_create_role(code, code) for code in sorted(desired)]
        target.roles = role_objs

        # 撤销负责人角色时，同事务清除其个人可选授权（不自动恢复）。
        revoked_optional = 0
        if (
            RoleCode.STUDENT_AFFAIRS_MANAGER.value in current_roles
            and RoleCode.STUDENT_AFFAIRS_MANAGER.value not in desired
        ):
            revoked_optional = self._repo.clear_optional_grants(target_user_id)

        target.lock_version += 1
        after_roles = sorted(desired)

        self._record_audit(
            actor_user_id=actor_user_id,
            action="role.assign",
            resource_type="user_account",
            resource_id=str(target_user_id),
            before={
                "roles": before_roles,
                "lock_version": expected_version,
                "cleared_optional_grants": revoked_optional,
            },
            after={
                "roles": after_roles,
                "lock_version": target.lock_version,
            },
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return after_roles, target.lock_version

    def set_optional_permission(
        self,
        *,
        actor_user_id: int,
        target_user_id: int,
        code: str,
        enabled: bool,
        expected_version: int,
        reason: str | None,
        request_id: str | None,
    ) -> tuple[str, bool, int]:
        """逐人开关学生工作负责人可选权限（三项许可清单，禁自我提权）。

        错误映射：非三项 code 422；目标非负责人 422；缺 optional_permission.manage
        或自我提权 403；目标不可见 404；版本不符 409。
        """
        locked = self._lock_accounts_ordered(actor_user_id, target_user_id)
        target = locked[target_user_id]

        # 功能权限 + 三项许可清单校验。
        self._require_actor_permission(
            actor_user_id, PermissionCode.OPTIONAL_PERMISSION_MANAGE.value
        )
        if code not in OPTIONAL_PERMISSION_CODES:
            raise AppError(
                ErrorCode.VALIDATION_ERROR,
                "仅可配置三项学生工作负责人可选权限",
                http_status=422,
                field_errors={"code": code},
            )

        # 禁止通过该接口给自己提权（PERMISSIONS.md 6.2）。
        if actor_user_id == target_user_id:
            raise PermissionDeniedError("不能为自己配置可选权限")

        # 目标必须当前具备负责人身份。
        target_roles = frozenset(self._repo.list_role_codes(target_user_id))
        if RoleCode.STUDENT_AFFAIRS_MANAGER.value not in target_roles:
            raise AppError(
                ErrorCode.VALIDATION_ERROR,
                "仅学生工作负责人可配置可选权限",
                http_status=422,
                field_errors={"target": "not_student_affairs_manager"},
            )

        if target.lock_version != expected_version:
            raise ConflictError(
                ErrorCode.VERSION_CONFLICT, "目标已被并发修改，请刷新后重试"
            )

        perm = self._repo.get_or_create_permission(code, code)
        if enabled:
            self._repo.set_user_permission(target_user_id, perm.id, actor_user_id)
        else:
            self._repo.revoke_user_permission(target_user_id, perm.id)

        target.lock_version += 1
        self._record_audit(
            actor_user_id=actor_user_id,
            action="optional_permission.update",
            resource_type="user_account",
            resource_id=str(target_user_id),
            before={
                "code": code,
                "enabled": not enabled,
                "lock_version": expected_version,
            },
            after={
                "code": code,
                "enabled": enabled,
                "lock_version": target.lock_version,
            },
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return code, enabled, target.lock_version

    # ---- 受限目标选择器（只读，不加锁、不提交）----
    def list_role_assignment_targets(
        self, actor_user_id: int
    ) -> tuple[list[RoleAssignmentTargetItem], list[str]]:
        actor_roles = frozenset(self._repo.list_role_codes(actor_user_id))
        assignable = assignable_roles(actor_roles)
        wechat_uids = set(self._session.scalars(select(WechatIdentity.user_id)).all())
        items = [
            RoleAssignmentTargetItem(
                id=str(user.id),
                username=user.username,
                display_name=user.display_name,
                status=user.status,
                roles=sorted(r.code for r in user.roles),
                lock_version=user.lock_version,
                student_id=str(user.student_id) if user.student_id is not None else None,
                has_wechat=user.id in wechat_uids,
            )
            for user in self._repo.list_all_users()
        ]
        return items, assignable

    def list_optional_permission_targets(
        self, actor_user_id: int
    ) -> tuple[list[OptionalPermissionTargetItem], list[str]]:
        configurable = sorted(OPTIONAL_PERMISSION_CODES)
        items: list[OptionalPermissionTargetItem] = []
        for user in self._repo.list_users_with_role(
            RoleCode.STUDENT_AFFAIRS_MANAGER.value
        ):
            granted = set(self._repo.list_optional_grant_codes(user.id))
            items.append(
                OptionalPermissionTargetItem(
                    id=str(user.id),
                    display_name=user.display_name,
                    status=user.status,
                    permissions=[
                        OptionalPermissionState(code=c, enabled=c in granted)
                        for c in configurable
                    ],
                    lock_version=user.lock_version,
                )
            )
        return items, configurable

    # ------------------------------------------------------------------ #
    # 绑定码管理与换绑（P2-D）：签发/作废一次性码、换绑撤销旧会话，均同事务审计
    # ------------------------------------------------------------------ #
    def issue_binding_token(
        self,
        *,
        actor_user_id: int,
        student_id: int,
        reason: str | None,
        request_id: str | None,
    ) -> tuple[int, str, datetime]:
        """为指定学生签发一次性 6 位绑定码；明文仅此返回一次，库中只存摘要。"""
        self._require_actor_permission(
            actor_user_id, PermissionCode.IDENTITY_BINDING_MANAGE.value
        )
        student = self._repo.get_student_by_id(student_id)
        if student is None:
            raise NotFoundError("学生不存在")

        code = generate_friendly_binding_code(6)
        expires_at = utcnow() + timedelta(
            minutes=self._settings.binding_token_valid_minutes
        )
        token = self._repo.create_binding_token(student.id, hash_token(code), expires_at)

        self._record_audit(
            actor_user_id=actor_user_id,
            action="binding_token.issue",
            resource_type="identity_binding_token",
            resource_id=str(token.id),
            before=None,
            after={
                "student_id": student.id,
                "expires_at": expires_at.isoformat(),
            },
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return token.id, code, expires_at

    def batch_issue_binding_tokens(
        self,
        *,
        actor_user_id: int,
        student_ids: list[int] | None = None,
        class_id: int | None = None,
        days_valid: int = 30,
        reason: str | None = None,
        request_id: str | None = None,
    ) -> list[dict]:
        """批量为学生签发 6 位绑定码。"""
        self._require_actor_permission(
            actor_user_id, PermissionCode.IDENTITY_BINDING_MANAGE.value
        )
        from app.modules.academic.models import AdministrativeClass, Student

        stmt = select(Student).where(Student.status == "ACTIVE")
        if student_ids:
            stmt = stmt.where(Student.id.in_(student_ids))
        elif class_id:
            stmt = stmt.where(Student.administrative_class_id == class_id)
        students = list(self._session.execute(stmt.order_by(Student.student_no)).scalars().all())

        bound_student_ids = {
            u.student_id
            for u in self._session.execute(
                select(UserAccount).where(UserAccount.student_id.is_not(None))
            ).scalars().all()
        }

        ac_ids = {s.administrative_class_id for s in students if s.administrative_class_id}
        acs = {
            ac.id: ac
            for ac in self._session.execute(
                select(AdministrativeClass).where(AdministrativeClass.id.in_(list(ac_ids)))
            ).scalars().all()
        } if ac_ids else {}

        results = []
        expires_at = utcnow() + timedelta(days=days_valid)

        for s in students:
            code = generate_friendly_binding_code(6)
            token = self._repo.create_binding_token(s.id, hash_token(code), expires_at)
            ac = acs.get(s.administrative_class_id)
            results.append({
                "student_id": s.id,
                "student_no": s.student_no,
                "name": s.name,
                "class_name": ac.class_name if ac else None,
                "college": ac.college if ac else None,
                "binding_code": code,
                "expires_at": expires_at.isoformat(),
                "is_bound": s.id in bound_student_ids,
            })

        self._record_audit(
            actor_user_id=actor_user_id,
            action="binding_token.batch_issue",
            resource_type="identity_binding_token",
            resource_id=f"batch:{len(results)}",
            before=None,
            after={"count": len(results), "expires_at": expires_at.isoformat()},
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return results

    def export_binding_tokens_excel(
        self,
        *,
        actor_user_id: int,
        student_ids: list[int] | None = None,
        class_id: int | None = None,
        days_valid: int = 30,
        reason: str | None = None,
        request_id: str | None = None,
    ):
        """批量生成并导出绑定码 Excel。"""
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        tokens_data = self.batch_issue_binding_tokens(
            actor_user_id=actor_user_id,
            student_ids=student_ids,
            class_id=class_id,
            days_valid=days_valid,
            reason=reason,
            request_id=request_id,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "学生绑定码"

        headers = ["学号", "姓名", "行政班级", "所属学院", "6位绑定码", "有效期截止", "绑定状态"]
        ws.append(headers)

        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1890FF", end_color="1890FF", fill_type="solid")
        cell_font = Font(name="微软雅黑", size=10)
        code_font = Font(name="Consolas", size=12, bold=True, color="D4380D")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        ws.row_dimensions[1].height = 28

        for row_idx, item in enumerate(tokens_data, start=2):
            status_text = "已绑定" if item["is_bound"] else "未绑定"
            ws.append([
                item["student_no"],
                item["name"],
                item["class_name"] or "",
                item["college"] or "",
                item["binding_code"],
                item["expires_at"][:19].replace("T", " "),
                status_text,
            ])
            ws.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx == 5:
                    cell.font = code_font
                    cell.alignment = center_align
                elif col_idx in (1, 6, 7):
                    cell.font = cell_font
                    cell.alignment = center_align
                else:
                    cell.font = cell_font
                    cell.alignment = left_align

        col_widths = [16, 12, 18, 18, 14, 22, 12]
        for idx, width in enumerate(col_widths, start=1):
            col_letter = chr(64 + idx)
            ws.column_dimensions[col_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def revoke_binding_token(
        self,
        *,
        actor_user_id: int,
        token_id: int,
        reason: str | None,
        request_id: str | None,
    ) -> int:
        """作废未使用的绑定码；已核销的码不可作废（409），已作废幂等返回。"""
        self._require_actor_permission(
            actor_user_id, PermissionCode.IDENTITY_BINDING_MANAGE.value
        )
        token = self._repo.get_binding_token_by_id(token_id)
        if token is None:
            raise NotFoundError("绑定码不存在")
        if token.status == BindingTokenStatus.USED.value:
            raise ConflictError(
                ErrorCode.STATE_CONFLICT, "绑定码已核销，无法作废"
            )
        if token.status == BindingTokenStatus.EXPIRED.value:
            return token_id  # 幂等：已作废无需重复

        token.status = BindingTokenStatus.EXPIRED.value
        self._record_audit(
            actor_user_id=actor_user_id,
            action="binding_token.revoke",
            resource_type="identity_binding_token",
            resource_id=str(token.id),
            before={"status": BindingTokenStatus.UNUSED.value},
            after={"status": BindingTokenStatus.EXPIRED.value},
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return token_id

    def student_binding_reset(
        self,
        *,
        actor_user_id: int,
        target_user_id: int,
        new_student_no: str | None,
        reason: str | None,
        request_id: str | None,
    ) -> tuple[int, int | None, int, int]:
        """换绑/解绑（线下核验后执行）：改绑或清除 student_id，并撤销旧会话。

        自动身份随绑定联动：新绑定确保持有 STUDENT；解绑移除 STUDENT。
        换绑后目标全部有效会话即时撤销，下一次请求需重新登录并读到最新身份。
        """
        locked = self._lock_accounts_ordered(actor_user_id, target_user_id)
        target = locked[target_user_id]
        self._require_actor_permission(
            actor_user_id, PermissionCode.IDENTITY_BINDING_MANAGE.value
        )

        before_student = target.student_id
        before_roles = self._repo.list_role_codes(target_user_id)

        new_student_id: int | None
        if new_student_no is not None:
            student = self._session.execute(select(Student).where(
                Student.student_no == new_student_no
            ).with_for_update().execution_options(populate_existing=True)).scalar_one_or_none()
            if student is None:
                raise NotFoundError("学号不存在")
            if student.status != "ACTIVE":
                raise ConflictError(ErrorCode.STATE_CONFLICT, "学生已停用，不可绑定")
            other = self._session.execute(
                select(UserAccount).where(
                    UserAccount.student_id == student.id,
                    UserAccount.id != target_user_id,
                )
            ).scalar_one_or_none()
            if other is not None:
                raise ConflictError(ErrorCode.STATE_CONFLICT, "该学生已被其他账号绑定")
            new_student_id = student.id
            target.student_id = student.id
            target.display_name = student.name
            student_role = self._repo.get_or_create_role(
                RoleCode.STUDENT.value, RoleCode.STUDENT.value
            )
            self._repo.grant_role(target_user_id, student_role.id)
        else:
            new_student_id = None
            target.student_id = None
            existing_student_role = self._repo.get_role_by_code(RoleCode.STUDENT.value)
            if existing_student_role is not None:
                self._repo.revoke_role(target_user_id, existing_student_role.id)

        if before_student != new_student_id or new_student_id is None:
            manager_role = self._repo.get_role_by_code(RoleCode.STUDENT_AFFAIRS_MANAGER.value)
            if manager_role is not None:
                self._repo.revoke_role(target_user_id, manager_role.id)
            self._repo.clear_optional_grants(target_user_id)

        revoked = self._repo.revoke_active_sessions(target_user_id, utcnow())
        target.lock_version += 1
        after_roles = self._repo.list_role_codes(target_user_id)

        self._record_audit(
            actor_user_id=actor_user_id,
            action="student.rebind",
            resource_type="user_account",
            resource_id=str(target_user_id),
            before={
                "student_id": before_student,
                "roles": before_roles,
                "lock_version": target.lock_version - 1,
            },
            after={
                "student_id": new_student_id,
                "display_name": target.display_name,
                "roles": after_roles,
                "lock_version": target.lock_version,
                "revoked_sessions": revoked,
            },
            reason=reason,
            request_id=request_id,
        )
        self._session.commit()
        return target_user_id, new_student_id, revoked, target.lock_version

    # ------------------------------------------------------------------ #
    # 学生绑定（原子：核销绑定码 + 写入绑定）
    # ------------------------------------------------------------------ #
    def bind_student(
        self,
        user_id: int,
        student_no: str,
        binding_code: str,
        *,
        request_id: str | None = None,
    ) -> tuple[int, str, str]:
        # 与管理换绑、停用保持账号→学生→绑定码顺序；锁后读取最新绑定状态。
        self._session.rollback()
        self._session.connection(execution_options={"isolation_level": "READ COMMITTED"})
        user = self._repo.get_user_by_id_for_update(user_id)
        if user is None:
            raise NotFoundError("账号不存在")
        if not user.is_active:
            raise UnauthenticatedError("账号已停用")
        # 尽早给出友好错误：已绑定者应走管理端换绑，而非再次自助核销。
        # （并发的"同一码双花"由下方持码锁后的状态检查兜住，此处非竞争点。）
        if user.student_id is not None:
            raise ConflictError(
                ErrorCode.STATE_CONFLICT,
                "该账号已绑定学生，如需换绑请走管理端流程",
            )

        student = self._session.execute(select(Student).where(
            Student.student_no == student_no
        ).with_for_update().execution_options(populate_existing=True)).scalar_one_or_none()
        if student is None:
            raise NotFoundError("学号不存在")
        if student.status != "ACTIVE":
            raise ConflictError(ErrorCode.STATE_CONFLICT, "学生已停用，不可绑定")

        # 一次性核销的串行点：对绑定码行加 FOR UPDATE 锁。并发核销同一码时，
        # 后到者持锁后读到先到者已提交的 USED 状态，从而被拒绝（保证不可双花）。
        # 账号与学生锁已在前面取得；RC 下可见其他账号刚提交的学生绑定。
        token = self._repo.get_binding_token_by_hash_for_update(hash_token(binding_code))
        if token is None:
            raise AppError(ErrorCode.VALIDATION_ERROR, "绑定码无效", http_status=422)

        now = utcnow()
        if token.status != BindingTokenStatus.UNUSED.value or token.used_at is not None:
            raise ConflictError(ErrorCode.STATE_CONFLICT, "绑定码已被使用")
        if token.expires_at <= now:
            token.status = BindingTokenStatus.EXPIRED.value
            self._session.commit()
            raise ConflictError(ErrorCode.STATE_CONFLICT, "绑定码已过期")
        if token.student_id != student.id:
            # 绑定码与提交学号不符：按无效处理，不泄露具体学生。
            token.failed_attempts += 1
            self._session.commit()
            raise AppError(ErrorCode.VALIDATION_ERROR, "绑定码与学号不匹配", http_status=422)

        # 持锁后的"已绑定"复查（读到并发事务已提交结果）：
        # 1) 本账号已绑定学生；2) 该学生已被其他账号绑定。均给出明确业务错误。
        if user.student_id is not None:
            raise ConflictError(
                ErrorCode.STATE_CONFLICT,
                "该账号已绑定学生，如需换绑请走管理端流程",
            )
        existing = self._session.execute(
            select(UserAccount).where(
                UserAccount.student_id == student.id, UserAccount.id != user.id
            )
        ).scalar_one_or_none()
        if existing is not None:
            existing_wechat = self._session.execute(
                select(WechatIdentity).where(WechatIdentity.user_id == existing.id)
            ).scalar_one_or_none()
            if existing_wechat is not None:
                raise ConflictError(ErrorCode.STATE_CONFLICT, "该学生已被其他账号绑定")
            # 预置系统账号关联：将当前微信登录身份与会话平滑关联至已有系统账号
            user_wechat = self._session.execute(
                select(WechatIdentity).where(WechatIdentity.user_id == user.id)
            ).scalar_one_or_none()
            if user_wechat is not None:
                user_wechat.user_id = existing.id
            self._session.execute(
                update(AuthSession).where(AuthSession.user_id == user.id).values(user_id=existing.id)
            )
            student_role = self._repo.get_or_create_role(
                RoleCode.STUDENT.value, RoleCode.STUDENT.value
            )
            self._repo.grant_role(existing.id, student_role.id)
            self._grant_volunteer_if_qualified(existing)
            token.status = BindingTokenStatus.USED.value
            token.used_at = now
            existing.display_name = student.name
            existing.lock_version += 1
            try:
                self._session.delete(user)
            except Exception:
                pass
            target_user = existing
        else:
            before_roles = self._repo.list_role_codes(user.id)
            user.student_id = student.id
            user.display_name = student.name
            token.status = BindingTokenStatus.USED.value
            token.used_at = now
            student_role = self._repo.get_or_create_role(
                RoleCode.STUDENT.value, RoleCode.STUDENT.value
            )
            self._repo.grant_role(user.id, student_role.id)
            user.lock_version += 1
            self._grant_volunteer_if_qualified(user)
            target_user = user

        after_roles = self._repo.list_role_codes(target_user.id)

        self._record_audit(
            actor_user_id=target_user.id,
            action="student.bind",
            resource_type="user_account",
            resource_id=str(target_user.id),
            before={"student_id": None},
            after={
                "student_id": student.id,
                "display_name": student.name,
                "roles": after_roles,
                "lock_version": target_user.lock_version,
            },
            reason=None,
            request_id=request_id,
        )
        self._session.commit()
        return student.id, student.student_no, student.name

    def _grant_volunteer_if_qualified(self, user: UserAccount) -> bool:
        """绑定后反向补授 VOLUNTEER：若该生已有 ACTIVE 学期下启用中的志愿者资格。

        与导入确认阶段的正向补授（importer._grant_volunteer_if_bound）对称，覆盖
        "先导入资格、后绑定微信"的时序。幂等：已持有 VOLUNTEER 则跳过。本方法不
        commit，交调用方在同一事务内提交。
        """
        if user.student_id is None:
            return False
        # 局部导入规避 identity↔academic 的模块级循环依赖风险。
        from app.modules.academic.models import Semester, VolunteerQualification

        has_qualified = self._session.execute(
            select(VolunteerQualification.id)
            .join(Semester, VolunteerQualification.semester_id == Semester.id)
            .where(
                VolunteerQualification.student_id == user.student_id,
                VolunteerQualification.enabled.is_(True),
                Semester.status == "ACTIVE",
            )
            .limit(1)
        ).scalar_one_or_none()
        if has_qualified is None:
            return False
        if RoleCode.VOLUNTEER.value in self._repo.list_role_codes(user.id):
            return False
        role = self._repo.get_or_create_role(
            RoleCode.VOLUNTEER.value, RoleCode.VOLUNTEER.value
        )
        self._repo.grant_role(user.id, role.id)
        user.lock_version += 1
        return True
