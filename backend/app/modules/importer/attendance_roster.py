"""学校原始课堂考勤表：首表为正式名单，辅助页仅作一致性核对。"""

import io
import re

from app.common.parsing.xlsx_tabular import _clean


def read_sheets(content: bytes) -> list[tuple[str, list[list[str]]]]:
    """保留工作表位置；限制与识别入口一致，避免旧预览绕过多表检查。"""
    if content.startswith(b"PK\x03\x04"):
        import zipfile

        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Excel 展开后超过 50 MiB，请拆分文件")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            if len(book.worksheets) > 20:
                raise ValueError("工作表超过 20 个，请只保留需要导入的资料")
            result = []
            for sheet in book:
                if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 256:
                    raise ValueError("表格超过 20000 行或 256 列，请拆分文件")
                result.append(
                    (
                        sheet.title,
                        [
                            [_clean(v) for v in row]
                            for row in sheet.iter_rows(
                                max_row=sheet.max_row or 20000,
                                max_col=sheet.max_column or 256,
                                values_only=True,
                            )
                        ],
                    )
                )
            return result
        finally:
            book.close()
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        import xlrd

        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if book.nsheets > 20:
                raise ValueError("工作表超过 20 个，请拆分文件")
            result = []
            for sheet in book.sheets():
                if sheet.nrows > 20000 or sheet.ncols > 256:
                    raise ValueError("表格超过 20000 行或 256 列，请拆分文件")
                result.append(
                    (
                        sheet.name,
                        [
                            [_clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                            for r in range(sheet.nrows)
                        ],
                    )
                )
            return result
        finally:
            book.release_resources()
    return []


def inspect_attendance(sheets: list[tuple[str, list[list[str]]]]) -> dict | None:
    """只识别有学院/班级标题及学号姓名表头的原表，不凭文件名猜测。"""
    if not sheets or not sheets[0][1]:
        return None
    name, matrix = sheets[0]
    title = " ".join(matrix[0])
    if not re.search(r"\S+学院\s+\S+\s+班级.*考勤表", title):
        return None
    header = next(
        (
            (i, row.index("学号"), row.index("姓名"))
            for i, row in enumerate(matrix[:15])
            if "学号" in row and "姓名" in row
        ),
        None,
    )
    if header is None:
        return None
    index, no_col, name_col = header
    primary = {}
    for row in matrix[index + 1 :]:
        if len(row) > max(no_col, name_col) and re.fullmatch(r"\d{6,20}", row[no_col]):
            primary[row[no_col]] = row[name_col]
    if not primary:
        return None
    issues, verified = [], []
    for sheet_name, rows in sheets[1:]:
        nonempty = [
            (i, [(j, value) for j, value in enumerate(row) if value])
            for i, row in enumerate(rows, 1)
            if any(row)
        ]
        if not nonempty:
            continue
        sheet_issues = []
        for row_no, cells in nonempty:
            # 辅助名单可位于任意相邻两列，但不得含未识别的额外业务列。
            if (
                len(cells) != 2
                or cells[1][0] != cells[0][0] + 1
                or not re.fullmatch(r"\d{6,20}", cells[0][1])
            ):
                code = "unsupported_auxiliary_sheet"
                message = "辅助页包含非名单内容，请核对后移出本文件，不会自动忽略"
            elif primary.get(cells[0][1]) != cells[1][1]:
                code = "auxiliary_roster_mismatch"
                message = "辅助页的学号/姓名与主表不一致或未在主表出现，请核对主表与辅助页"
            else:
                continue
            sheet_issues.append(
                {
                    "code": code,
                    "message": f"{sheet_name} 第 {row_no} 行：{message}",
                    "field": "学号/姓名",
                    "row": row_no,
                    "sheet": sheet_name,
                    "severity": "error",
                }
            )
        issues.extend(sheet_issues)
        if not sheet_issues:
            verified.append(sheet_name)
    return {"source_sheet": name, "verified_auxiliary_sheets": verified, "issues": issues}
