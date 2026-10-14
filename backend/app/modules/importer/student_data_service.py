"""教师数据交接：可回读导出、范围内差异计划、版本复核、整批确认。

不根据文件缺行自动删除，也不根据旧源班级覆盖已确认异动。
"""

import io
import zipfile
from datetime import date, timedelta
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import or_, select, true
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.exceptions import AppError, ConflictError, ErrorCode, NotFoundError
from app.modules.academic.enrollment_models import College, StudentDataPlan, StudentEnrollment
from app.modules.academic.enrollment_schemas import EnrollmentChangeRequest
from app.modules.academic.enrollment_service import EnrollmentService, class_snapshot
from app.modules.academic.models import AdministrativeClass, Student
from app.modules.audit.models import AuditLog
from app.modules.identity.authorization import business_today, require_resource_scope, resolve_scope
from app.modules.identity.models import UserAccount
from app.modules.identity.repository import IdentityRepository


class StudentDataRow(BaseModel):
    model_config = ConfigDict(extra="ignore", str_strip_whitespace=True)
    student_no: str = Field(min_length=1, max_length=32)
    name: str = Field(min_length=1, max_length=128)
    class_code: str = Field(min_length=1, max_length=64)
    lock_version: int | None = Field(default=None, ge=0)
    source_as_of: date | None = None
    kind: Literal["TRANSFER", "RETURN", "SUSPEND", "GRADUATE", "WITHDRAW"] | None = None
    effective_date: date | None = None


class StudentDataPreviewRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    college_id: int = Field(ge=1)
    purpose: Literal["INCREMENTAL", "FULL_COMPARE", "ENROLLMENT_CHANGE"] = "INCREMENTAL"
    rows: list[StudentDataRow] = Field(max_length=5000)


class StudentDataService:
    def __init__(self, db: Session):
        self.db = db
        self.today = business_today()

    def _scope(self, actor_id: int, college_id: int, action: str):
        scope = resolve_scope(self.db, actor_id, action)
        require_resource_scope(scope, college_id)
        college = self.db.get(College, college_id)
        if college is None or college.status != "ACTIVE":
            raise NotFoundError("学院不存在或已停用")
        return scope

    def _students(self, college_id: int) -> list[Student]:
        current = select(StudentEnrollment.student_id).where(
            StudentEnrollment.college_id == college_id,
            StudentEnrollment.superseded.is_(False),
            StudentEnrollment.effective_from <= self.today,
            or_(
                StudentEnrollment.effective_to.is_(None),
                StudentEnrollment.effective_to > self.today,
            ),
        )
        any_timeline = select(StudentEnrollment.student_id)
        classes = select(AdministrativeClass.id).where(AdministrativeClass.college_id == college_id)
        return list(
            self.db.scalars(
                select(Student)
                .where(
                    or_(
                        Student.id.in_(current),
                        Student.id.not_in(any_timeline)
                        & Student.administrative_class_id.in_(classes),
                    )
                )
                .order_by(Student.student_no)
            )
        )

    def preview(
        self, actor_id: int, *, college_id: int, rows: list[dict], purpose: str = "INCREMENTAL"
    ) -> dict[str, Any]:
        scope = self._scope(actor_id, college_id, "student.manage")
        request = StudentDataPreviewRequest(college_id=college_id, rows=rows, purpose=purpose)
        classes = {
            ac.class_code: ac
            for ac in self.db.scalars(
                select(AdministrativeClass).where(
                    true()
                    if scope.kind == "SCHOOL" and purpose == "ENROLLMENT_CHANGE"
                    else AdministrativeClass.college_id == college_id,
                    AdministrativeClass.status == "ACTIVE",
                )
            )
        }
        incoming = set()
        normalized_numbers, seen_student_ids = set(), set()
        result, errors = [], []
        timeline = EnrollmentService(self.db)
        for index, item in enumerate(request.rows, start=1):
            if item.student_no.casefold() in normalized_numbers:
                errors.append({"row": index, "code": "DUPLICATE_STUDENT_NO"})
                continue
            incoming.add(item.student_no)
            normalized_numbers.add(item.student_no.casefold())
            ac = classes.get(item.class_code)
            if ac is None:
                errors.append({"row": index, "code": "CLASS_NOT_IN_WORKSPACE"})
                continue
            student = self.db.scalar(select(Student).where(Student.student_no == item.student_no))
            if (
                student is None
                and self.db.scalar(
                    select(Student.__table__.c.id).where(
                        Student.__table__.c.student_no == item.student_no
                    )
                )
                is not None
            ):
                errors.append({"row": index, "code": "STUDENT_OUTSIDE_WORKSPACE"})
                continue
            enrollment = timeline.at(student.id, self.today) if student else None
            if student and student.id in seen_student_ids:
                errors.append({"row": index, "code": "DUPLICATE_STUDENT_IDENTITY"})
                continue
            if student:
                seen_student_ids.add(student.id)
            action = "CREATE"
            if student:
                current_ac = (
                    self.db.get(
                        AdministrativeClass,
                        enrollment.administrative_class_id
                        if enrollment
                        else student.administrative_class_id,
                    )
                    if (
                        enrollment
                        and enrollment.administrative_class_id
                        or student.administrative_class_id
                    )
                    else None
                )
                current_college = (
                    enrollment.college_id
                    if enrollment
                    else current_ac.college_id
                    if current_ac
                    else None
                )
                try:
                    require_resource_scope(scope, current_college)
                except AppError:
                    errors.append({"row": index, "code": "STUDENT_OUTSIDE_WORKSPACE"})
                    continue
                if purpose == "ENROLLMENT_CHANGE":
                    if item.lock_version is not None and item.lock_version != student.lock_version:
                        errors.append({"row": index, "code": "STALE_SOURCE"})
                        continue
                    if not item.kind or not item.effective_date:
                        errors.append({"row": index, "code": "CHANGE_DATE_AND_KIND_REQUIRED"})
                        continue
                    command = EnrollmentChangeRequest(
                        kind=item.kind,
                        effective_date=item.effective_date,
                        target_class_id=ac.id if item.kind in {"TRANSFER", "RETURN"} else None,
                        lock_version=student.lock_version,
                    )
                    try:
                        timeline.preview(actor_id, student.id, command)
                    except AppError as exc:
                        errors.append(
                            {"row": index, "code": exc.code.value, "message": exc.message}
                        )
                        continue
                    action = "ENROLLMENT_CHANGE"
                elif item.lock_version is not None and item.lock_version != student.lock_version:
                    action = "STALE_SOURCE"
                elif student.status != "ACTIVE" or enrollment and enrollment.status != "ENROLLED":
                    action = "LIFECYCLE_REQUIRES_CHANGE"
                elif (
                    enrollment.administrative_class_id
                    if enrollment
                    else student.administrative_class_id
                ) != ac.id:
                    action = "TRANSFER_REQUIRES_CHANGE"
                elif student.name != item.name:
                    action = "UPDATE_NAME"
                else:
                    action = "UNCHANGED"
            elif purpose == "ENROLLMENT_CHANGE":
                errors.append({"row": index, "code": "STUDENT_NOT_FOUND"})
                continue
            result.append(
                {
                    "row": index,
                    "student_id": str(student.id) if student else None,
                    "student_no": item.student_no,
                    "name": item.name,
                    "class_id": str(ac.id),
                    "class_code": ac.class_code,
                    "action": action,
                    "baseline_version": student.lock_version if student else None,
                    "class_version": ac.updated_at.isoformat(),
                    "source": item.model_dump(mode="json"),
                }
            )
        missing_students = (
            [
                s
                for s in self._students(college_id)
                if s.student_no not in incoming
                and (
                    (e := timeline.at(s.id, self.today)) is not None
                    and e.status == "ENROLLED"
                    or e is None
                    and s.status == "ACTIVE"
                )
            ]
            if purpose == "FULL_COMPARE"
            else []
        )
        missing = [str(s.id) for s in missing_students]
        plan = StudentDataPlan(
            actor_user_id=actor_id,
            college_id=college_id,
            purpose=purpose,
            expires_at=self.today + timedelta(days=1),
            payload={
                "rows": result,
                "errors": errors,
                "missing_student_ids": missing,
                "missing_students": [
                    {"id": str(s.id), "student_no": s.student_no, "name": s.name}
                    for s in missing_students
                ],
                "scope_baseline": {str(s.id): s.lock_version for s in self._students(college_id)}
                if purpose == "FULL_COMPARE"
                else {},
            },
        )
        self.db.add(plan)
        self.db.commit()
        return self._view(plan)

    @staticmethod
    def _view(plan: StudentDataPlan) -> dict[str, Any]:
        return {
            "id": str(plan.id),
            "college_id": str(plan.college_id),
            "purpose": plan.purpose,
            "status": plan.status,
            "lock_version": plan.lock_version,
            "expires_at": plan.expires_at.isoformat(),
            **plan.payload,
            "can_confirm": not plan.payload.get("errors"),
            "missing_policy": "REVIEW_ONLY",
        }

    def confirm(
        self, actor_id: int, plan_id: int, expected_version: int, *, request_id: str | None = None
    ) -> dict[str, Any]:
        self.db.rollback()
        plan = self.db.get(StudentDataPlan, plan_id)
        if plan is None:
            raise NotFoundError("差异计划不存在")
        scope = self._scope(actor_id, plan.college_id, "student.manage")
        ids = sorted(int(r["student_id"]) for r in plan.payload["rows"] if r["student_id"])
        user_ids = set(
            self.db.scalars(select(UserAccount.id).where(UserAccount.student_id.in_(ids)))
        ) | {actor_id}
        from app.modules.academic.lifecycle_transaction import begin_lifecycle_write

        begin_lifecycle_write(self.db)  # 释放预检操作者锁，再按账号 ID 顺序加锁。
        repo = IdentityRepository(self.db)
        for uid in sorted(user_ids):
            repo.get_user_by_id_for_update(uid)
        scope = self._scope(actor_id, plan.college_id, "student.manage")
        EnrollmentService(self.db)._lock_semesters()
        self.db.refresh(plan, with_for_update=True)
        if plan.lock_version != expected_version or plan.status != "PREVIEW":
            raise ConflictError(ErrorCode.VERSION_CONFLICT, "计划已确认或版本已变化")
        if plan.expires_at <= self.today or plan.payload.get("errors"):
            raise ConflictError(message="计划过期或有错误，请重新预览")
        class_ids = sorted({int(r["class_id"]) for r in plan.payload["rows"]})
        self.db.execute(
            select(AdministrativeClass)
            .where(AdministrativeClass.id.in_(class_ids))
            .order_by(AdministrativeClass.id)
            .with_for_update()
        ).all()
        locked = {
            s.id: s
            for s in self.db.scalars(
                select(Student)
                .where(Student.id.in_(ids))
                .order_by(Student.id)
                .with_for_update()
                .execution_options(populate_existing=True)
            )
        }
        if (
            plan.purpose == "FULL_COMPARE"
            and {str(s.id): s.lock_version for s in self._students(plan.college_id)}
            != plan.payload["scope_baseline"]
        ):
            raise ConflictError(message="学院学生范围已变化，请重新比较")
        for row in plan.payload["rows"]:
            student = locked.get(int(row["student_id"])) if row["student_id"] else None
            ac = self.db.get(AdministrativeClass, int(row["class_id"]), populate_existing=True)
            if (
                ac is None
                or (ac.college_id != plan.college_id and scope.kind != "SCHOOL")
                or ac.status != "ACTIVE"
                or ac.updated_at.isoformat() != row["class_version"]
            ):
                raise ConflictError(message="班级归属或资料已变化，请重新预览")
            if row["student_id"] and (
                student is None or student.lock_version != row["baseline_version"]
            ):
                raise ConflictError(ErrorCode.VERSION_CONFLICT, "学生在预览后已修改，整批未应用")
            if (
                not row["student_id"]
                and self.db.scalar(
                    select(Student.id).where(Student.student_no == row["student_no"])
                )
                is not None
            ):
                raise ConflictError(message="学号已存在，请重新预览")
        applied = []
        changed_students: set[int] = set()
        refresh_dates = []
        # 所有基线先核实，再按学生 ID 顺序落库；任一失败整批回滚。
        try:
            for row in sorted(
                plan.payload["rows"],
                key=lambda r: (int(r["student_id"]) if r["student_id"] else 2**64, r["student_no"]),
            ):
                student = locked.get(int(row["student_id"])) if row["student_id"] else None
                ac = self.db.get(AdministrativeClass, int(row["class_id"]))
                if ac is None:
                    raise ConflictError(message="班级不存在")
                if row["action"] == "CREATE":
                    student = Student(
                        student_no=row["student_no"],
                        name=row["name"],
                        administrative_class_id=ac.id,
                    )
                    self.db.add(student)
                    self.db.flush()
                    self.db.add(
                        StudentEnrollment(
                            student_id=student.id,
                            administrative_class_id=ac.id,
                            college_id=ac.college_id,
                            effective_from=self.today,
                            status="ENROLLED",
                            revision=0,
                            historical_confirmed=False,
                            snapshot=class_snapshot(ac),
                        )
                    )
                elif row["action"] == "UPDATE_NAME":
                    if student is None:
                        raise ConflictError(message="学生不存在")
                    student.name = row["name"]
                    student.lock_version += 1
                    self.db.flush()
                elif row["action"] == "ENROLLMENT_CHANGE":
                    if student is None:
                        raise ConflictError(message="学生不存在")
                    item = StudentDataRow.model_validate(row["source"])
                    if item.kind is None or item.effective_date is None:
                        raise ConflictError(message="计划缺少异动类型或日期")
                    command = EnrollmentChangeRequest(
                        kind=item.kind,
                        effective_date=item.effective_date,
                        target_class_id=ac.id if item.kind in {"TRANSFER", "RETURN"} else None,
                        lock_version=student.lock_version,
                        idempotency_key=f"student-plan-{plan.id}-{student.id}",
                    )
                    EnrollmentService(self.db).commit(
                        actor_id,
                        student.id,
                        command,
                        request_id=request_id,
                        _owns_transaction=False,
                        _refresh_task_state=False,
                    )
                    refresh_dates.append(item.effective_date)
                else:
                    continue
                if student is not None:
                    changed_students.add(student.id)
                    if row["action"] != "ENROLLMENT_CHANGE":
                        refresh_dates.append(self.today)
                    applied.append({"student_id": str(student.id), "action": row["action"]})
            self.db.flush()
            refresh = (
                EnrollmentService(self.db)._refresh_tasks(
                    0, min(refresh_dates), actor_id, student_ids=changed_students
                )
                if changed_students
                else {}
            )
            plan.status = "CONFIRMED"
            plan.lock_version += 1
            plan.payload = {**plan.payload, "applied": applied, "task_refresh": refresh}
            self.db.add(
                AuditLog(
                    actor_user_id=actor_id,
                    action="import.student_data.confirm",
                    resource_type="student_data_plan",
                    resource_id=str(plan.id),
                    before_json={"status": "PREVIEW"},
                    after_json={
                        "status": "CONFIRMED",
                        "college_id": str(plan.college_id),
                        "purpose": plan.purpose,
                        "applied": applied,
                    },
                    request_id=request_id,
                )
            )
            self.db.commit()
        except IntegrityError as error:
            self.db.rollback()
            raise ConflictError(
                message="唯一编码或关联在预览后发生变化，整批未应用，请重新预览"
            ) from error
        except Exception:
            self.db.rollback()
            raise
        return self._view(plan)

    def export(self, actor_id: int, college_id: int, *, include_archive: bool = False) -> bytes:
        self._scope(actor_id, college_id, "student.read")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "学生交接"
        headers = [
            "student_id",
            "student_no",
            "name",
            "class_code",
            "class_name",
            "college_code",
            "status",
            "lock_version",
            "source_as_of",
            "history_until",
        ]
        sheet.append(headers)
        college = self.db.get(College, college_id)
        if college is None:
            raise NotFoundError("学院不存在")
        timeline = EnrollmentService(self.db)
        for student in self._students(college_id):
            e = timeline.at(student.id, self.today)
            status = e.status if e else "ENROLLED" if student.status == "ACTIVE" else student.status
            if not include_archive and status != "ENROLLED":
                continue
            class_id = e.administrative_class_id if e else student.administrative_class_id
            ac = self.db.get(AdministrativeClass, class_id) if class_id else None
            sheet.append(
                [
                    str(student.id),
                    student.student_no,
                    student.name,
                    ac.class_code if ac else "",
                    e.snapshot.get("class_name") if e else ac.class_name if ac else "",
                    college.code,
                    status,
                    student.lock_version,
                    self.today.isoformat(),
                    student.history_until.isoformat() if student.history_until else "",
                ]
            )
            # 学号和外部字符串强制文本，避免丢前导零或公式注入。
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def preview_file(
        self, actor_id: int, college_id: int, content: bytes, *, purpose: str = "INCREMENTAL"
    ) -> dict[str, Any]:
        self._scope(actor_id, college_id, "student.manage")
        return self.preview(
            actor_id, college_id=college_id, rows=self._file_rows(content), purpose=purpose
        )

    def _file_rows(self, content: bytes, *, allow_bundle: bool = True) -> list[dict]:
        if len(content) > 10 * 1024 * 1024:
            raise AppError(ErrorCode.VALIDATION_ERROR, "文件超过 10 MiB", http_status=422)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 200 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                    raise ValueError("文件解压大小超限")
                if "[Content_Types].xml" not in archive.namelist():
                    if not allow_bundle:
                        raise ValueError("批量文件不允许嵌套归档")
                    files = [e for e in entries if not e.is_dir()]
                    if not files or any(not e.filename.lower().endswith(".xlsx") for e in files):
                        raise ValueError("批量 ZIP 仅接受 xlsx")
                    rows = []
                    for entry in files:
                        rows.extend(self._file_rows(archive.read(entry), allow_bundle=False))
                    if len(rows) > 5000:
                        raise ValueError("整批最多 5000 人")
                    return rows
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            if len(workbook.sheetnames) != 1:
                raise ValueError("每个文件仅允许一个数据表，避免静默漏导")
            sheet = workbook.active
            if sheet is None or sheet.max_column > 32:
                workbook.close()
                raise ValueError("表格列数超限")
            iterator = sheet.values
            headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
            if not {"student_no", "name", "class_code"} <= set(headers) or len(set(headers)) != len(
                headers
            ):
                raise ValueError("模板须含唯一 student_no/name/class_code 表头")
            rows = []
            for values in iterator:
                if not any(v is not None for v in values):
                    continue
                row = {
                    k: v
                    for k, v in zip(headers, values, strict=False)
                    if k and v is not None and v != ""
                }
                if not isinstance(row.get("student_no"), str):
                    raise ValueError("学号必须使用文本单元格，避免丢失前导零")
                rows.append(row)
                if len(rows) > 5000:
                    raise ValueError("单批最多 5000 人")
            workbook.close()
        except Exception as exc:
            raise AppError(
                ErrorCode.VALIDATION_ERROR,
                "学生交接文件格式无效，请使用导出模板并核对文本学号",
                http_status=422,
            ) from exc
        return rows
