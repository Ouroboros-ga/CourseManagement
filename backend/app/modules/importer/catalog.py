"""资料目录及模板包；required 表示准备建议，不构成全局操作门禁。"""

from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.modules.importer.permissions import (
    STAFF_IMPORT_ACTION,
    TARGET_MANAGE_PERMISSION,
    template_spec,
)
from app.modules.importer.staff_adapter import get_template

CATALOG = (
    ("admin_roster", "行政班与学生名单", True),
    ("faculty_staff", "教职工档案", True),
    ("class_responsibility", "班主任与辅导员责任关系", True),
    ("timetable", "本学期标准课表", True),
    ("grid_timetable", "教务网格课表", True),
    ("volunteer", "志愿者学期资格底册", True),
    ("roster", "教学班补充名单", False),
    ("elective_course", "选修与分层课补充资料", False),
)
_STAFF_TARGETS = {"faculty_staff", "class_responsibility"}


def _spec(target: str) -> dict[str, Any]:
    if target in _STAFF_TARGETS:
        return get_template(target).model_dump()
    return template_spec(target)


def catalog_items() -> list[dict[str, Any]]:
    """标准课表与网格课表二选一；已有基础资料可直接复用。"""
    return [
        {
            "target": target,
            "label": label,
            "required": required,
            "semester_required": "semester_id" in _spec(target)["scope_fields"],
            "manage_permission": (
                STAFF_IMPORT_ACTION if target in _STAFF_TARGETS
                else TARGET_MANAGE_PERMISSION[target]
            ),
        }
        for target, label, required in CATALOG
    ]


def _normalize(value: str) -> str:
    return value.strip().lstrip("\ufeff").casefold()


def detect_target(headers: list[str], filename: str = "") -> str | None:
    """仅表头证据充分且唯一时识别；文件名不能授予资料类型。"""
    normalized = {_normalize(header) for header in headers if header.strip()}
    candidates: dict[str, set[str]] = {}
    for target, _, _ in CATALOG:
        columns = _spec(target)["columns"]
        matched: set[str] = set()
        responsibility_evidence = False
        complete = True
        for column in columns:
            aliases = {
                _normalize(str(value))
                for value in [column["header"], column["key"], *column["aliases"]]
            }
            found = aliases & normalized
            if found and column["key"].startswith(("head_teacher", "counselor")):
                responsibility_evidence = True
            matched.update(found)
            if column["required"] and not found:
                complete = False
        # 责任表仅“行政班”一列虽满足字段必填约束，却不足以辨识资料用途。
        if target == "class_responsibility" and not responsibility_evidence:
            complete = False
        if complete and matched:
            candidates[target] = matched
    # 简单名单和志愿者只有学号时不可区分；更完整的表头可排除其子集。
    winners = [
        target for target, matched in candidates.items()
        if not any(matched < other for other in candidates.values())
    ]
    return winners[0] if len(winners) == 1 else None


def build_template_pack(targets: list[str]) -> bytes:
    """按现有字段契约生成独立模板；不伪造学校网格课表格式。"""
    labels = {target: label for target, label, _ in CATALOG}
    selected = list(dict.fromkeys(targets))
    unknown = set(selected) - labels.keys()
    if unknown:
        raise ValueError(f"不支持的资料类型：{', '.join(sorted(unknown))}")
    output = BytesIO()
    notes = [
        "学院资料模板说明",
        "已有基础档案无需重复上传。各类资料独立预览、核对与保存。",
        "标准课表与教务网格课表二选一；教学班补充名单及选修资料按需使用。",
        "模板第二行起为示例，请替换为真实资料；学号、工号等请保持文本格式。",
        "姓名仅用于展示与辅助核对，人员关系以学号或工号为准。",
    ]
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for target in selected:
            spec = _spec(target)
            if target == "grid_timetable":
                notes.append("教务网格课表：直接上传已支持的学校教务原始 .xls/.xlsx 文件。")
                continue
            columns = spec["columns"]
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.title = "资料"
            sheet.append([column["header"] for column in columns])
            for row in spec["example_rows"]:
                sheet.append([row.get(column["header"], "") for column in columns])
            for row in sheet:
                for cell in row:
                    cell.number_format = "@"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EFF6")
                sheet.column_dimensions[cell.column_letter].width = 24
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            stream = BytesIO()
            workbook.save(stream)
            archive.writestr(f"{labels[target]}模板.xlsx", stream.getvalue())
            required = "、".join(
                column["header"] for column in columns if column["required"]
            )
            notes.append(f"{labels[target]}：必填列 {required}。")
            if spec["scope_fields"]:
                notes.append(f"  导入时需选择：{', '.join(spec['scope_fields'])}。")
        archive.writestr("导入说明.txt", "\n".join(notes).encode("utf-8-sig"))
    return output.getvalue()
