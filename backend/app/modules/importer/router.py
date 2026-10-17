"""导入路由：multipart 上传→预览、确认、错误、模板（技术方案 19，PERMISSIONS.md 12.5）。

前缀在 main.py 挂载为 /api/v1。功能守卫：导入相关端点要求 `import.execute`
（路由级早拦），目标资源 manage 与组合权限在 Service 事务内纵深复核；模板下载
仅需目标 manage（Service 内动态判定，路由仅要求已登录）。
"""

from __future__ import annotations

from typing import Annotated

from fastapi import APIRouter, Depends, File, Form, Query, Request, UploadFile
from sqlalchemy.orm import Session

from app.common.responses import success
from app.core.database import get_db
from app.core.permissions import PermissionCode
from app.modules.identity.deps import CurrentUserDep, require_permission
from app.modules.importer.service import ImporterService

router = APIRouter(tags=["import"])

ImportExecuteDep = Annotated[
    CurrentUserDep, Depends(require_permission(PermissionCode.IMPORT_EXECUTE.value))
]


def get_service(request: Request, actor: CurrentUserDep,
                db: Annotated[Session, Depends(get_db)]) -> ImporterService:
    from app.core.config import get_settings
    from app.core.exceptions import AppError, ErrorCode
    from app.modules.importer.scoped_service import ScopedImporterService

    cid = request.path_params.get("cid")
    if cid is not None:
        if not cid.isascii() or not cid.isdecimal() or not 0 < int(cid) < 2**63:
            raise AppError(ErrorCode.VALIDATION_ERROR, "学院 ID 不合法", 422)
        return ScopedImporterService(db, actor, int(cid))
    if get_settings().campus_business_scoped_enabled:
        raise AppError(ErrorCode.STATE_CONFLICT, "请从已授权学院工作区执行导入", 409)
    return ImporterService(db)


ServiceDep = Annotated[ImporterService, Depends(get_service)]


def _rid(request: Request) -> str | None:
    return getattr(request.state, "request_id", None)


# ---- 模板（仅需目标 manage，路由仅要求登录，具体在 Service 判定）----
@router.get("/import-templates/{target}")
def get_import_template(
    target: str, actor: CurrentUserDep, service: ServiceDep, request: Request
) -> dict[str, object]:
    result = service.get_template(actor, target)
    return success(result.model_dump(), _rid(request))


@router.get("/import-templates/{target}/file")
def download_import_template(
    target: str, actor: CurrentUserDep, service: ServiceDep, request: Request
):
    import io
    import urllib.parse

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from starlette.responses import Response

    result = service.get_template(actor, target)
    spec = result.model_dump()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "导入数据"

    columns = spec.get("columns", [])
    headers = [col["header"] for col in columns]
    if headers:
        ws.append(headers)

    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1F2937")
    header_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    data_font = Font(name="Microsoft YaHei", size=10)
    for row_data in spec.get("example_rows", []):
        row_values = []
        for col in columns:
            row_values.append(row_data.get(col["header"], ""))
        ws.append(row_values)
        for cell in ws[ws.max_row]:
            cell.font = data_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len * 2 + 6, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    target_names = {
        "admin_roster": "行政班与学生名单模板",
        "class_duty": "班级管理责任关系模板",
        "timetable": "学院教务课表标准长表模板",
        "volunteer": "志愿者学期资格底册模板",
        "roster": "教学班补充名单模板",
        "grid_timetable": "学校网格课表说明",
    }
    filename = f"{target_names.get(target, target + '_template')}.xlsx"
    encoded_filename = urllib.parse.quote(filename)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


# ---- 创建预览（上传 xlsx，解析暂存，不写业务表）----
@router.post("/imports")
async def create_import(
    request: Request,
    actor: ImportExecuteDep,
    service: ServiceDep,
    file: Annotated[UploadFile, File()],
    target: Annotated[str, Form()],
    semester_id: Annotated[int | None, Form()] = None,
    teaching_class_id: Annotated[int | None, Form()] = None,
    course_type_id: Annotated[int | None, Form()] = None,
    replace: Annotated[bool, Form()] = True,
) -> dict[str, object]:
    from app.core.exceptions import AppError, ErrorCode
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise AppError(ErrorCode.VALIDATION_ERROR, "单个导入文件不能超过 10 MiB", 422)
    resolved_target = "admin_roster" if target == "class_duty" else target
    result = service.create_preview(
        actor,
        target=resolved_target,
        content=content,
        semester_id=semester_id,
        teaching_class_id=teaching_class_id,
        course_type_id=course_type_id,
        filename=file.filename,
        replace=replace,
        request_id=_rid(request),
    )
    return success(result.model_dump(), _rid(request))


# ---- 最近导入列表（学院工作区资料准备首页；旧兼容入口由 Service 显式拒绝）----
@router.get("/imports")
def list_imports(
    actor: ImportExecuteDep,
    service: ServiceDep,
    request: Request,
    limit: Annotated[int, Query(ge=1, le=50)] = 10,
) -> dict[str, object]:
    result = service.list_recent(actor, limit)
    return success([item.model_dump() for item in result], _rid(request))


@router.get("/imports/{batch_id}")
def get_import(
    batch_id: int, actor: ImportExecuteDep, service: ServiceDep, request: Request
) -> dict[str, object]:
    result = service.get_batch(actor, batch_id)
    return success(result.model_dump(), _rid(request))


@router.post("/imports/{batch_id}/confirm")
def confirm_import(
    batch_id: int, actor: ImportExecuteDep, service: ServiceDep, request: Request
) -> dict[str, object]:
    result = service.confirm(actor, batch_id, _rid(request))
    return success(result.model_dump(), _rid(request))


@router.get("/imports/{batch_id}/errors")
def import_errors(
    batch_id: int,
    actor: ImportExecuteDep,
    service: ServiceDep,
    request: Request,
    _q: Annotated[str | None, Query()] = None,  # 预留过滤参数位
) -> dict[str, object]:
    result = service.get_errors(actor, batch_id)
    return success(result, _rid(request))


@router.post("/imports/bulk")
async def bulk_import(
    request: Request,
    actor: ImportExecuteDep,
    service: ServiceDep,
    semester_id: Annotated[int, Form()],
    files: Annotated[list[UploadFile], File()],
    course_type_id: Annotated[int | None, Form()] = None,
) -> dict[str, object]:
    """旧兼容入口逐文件预览并确认；前面的成功文件不会随后续文件失败回滚。"""
    raw_files: list[tuple[str, bytes]] = []
    for f in files:
        data = await f.read()
        raw_files.append((f.filename or "file", data))
    result = service.import_bulk_files(
        actor,
        semester_id=semester_id,
        files=raw_files,
        course_type_id=course_type_id,
        request_id=_rid(request),
    )
    return success(result, _rid(request))
