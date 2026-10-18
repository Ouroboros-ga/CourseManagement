"""无持久化的文件识别；正式预览/保存仍使用原导入接口。"""

import io
import zipfile

from app.common.parsing.xlsx_tabular import read_any_matrix
from app.core.exceptions import AppError, ErrorCode
from app.modules.importer.catalog import detect_target
from app.modules.importer.limits import MAX_ARCHIVE_BYTES
from app.modules.importer.zip_preview import MAX_FILE, expand_zip


def _inspect(content: bytes, name: str) -> dict:
    result = {"filename": name, "target": None, "blocked": False, "message": ""}
    try:
        # XLSX 本身也是压缩包；先限制展开体积，再让表格解析器读取。
        if zipfile.is_zipfile(io.BytesIO(content)):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("Excel 展开后超过 50 MiB，请拆分文件")
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                if len(book.worksheets) > 20:
                    raise ValueError("工作表超过 20 个，请只保留需要导入的资料")
                populated = []
                for sheet in book.worksheets:
                    if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 256:
                        raise ValueError("表格超过 20000 行或 256 列，请拆分文件")
                    if any(
                        any(v is not None for v in row)
                        for row in sheet.iter_rows(
                            max_row=sheet.max_row or 20000,
                            max_col=sheet.max_column or 1,
                            values_only=True,
                        )
                    ):
                        populated.append(sheet.title)
                if len(populated) > 1:
                    raise ValueError("文件包含多个有内容的工作表，请拆成单独文件后导入，避免遗漏")
                if populated and populated[0] != book.worksheets[0].title:
                    raise ValueError("数据不在第一个工作表，请将数据单独另存为一个文件")
            finally:
                book.close()
        elif content.startswith(b"\xd0\xcf\x11\xe0"):
            import xlrd

            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                if sum(bool(s.nrows and s.ncols) for s in book.sheets()) > 1:
                    raise ValueError("文件包含多个工作表，请拆成单独文件后导入")
                if any(s.nrows > 20000 or s.ncols > 256 for s in book.sheets()):
                    raise ValueError("表格超过 20000 行或 256 列，请拆分文件")
            finally:
                book.release_resources()
        matrix = read_any_matrix(content)
        if len(matrix) > 20000 or any(len(row) > 256 for row in matrix):
            raise ValueError("表格超过 20000 行或 256 列，请拆分文件")
        # 标准长表要求首行表头；旧花名册、网格原表另由既有格式解析器识别。
        target = detect_target([str(v) for v in matrix[0]], name) if matrix else None
        if target is None:
            from app.common.parsing.grid import load_grid

            try:
                class_name, rows = load_grid(content)
                if rows and any(r.get("course_name") for r in rows):
                    target = "grid_timetable"
            except Exception:  # 非网格格式不是文件损坏，保留给用户选择原解析器。
                pass
        result["target"] = target
        if target is None:
            result["message"] = "未能唯一识别资料类型，请选择；读取预览后仍会校验文件内容"
    except Exception as exc:
        result.update(blocked=True, message=f"无法读取：{exc}")
    return result


def inspect_upload(content: bytes, filename: str) -> list[dict]:
    limit = MAX_ARCHIVE_BYTES if filename.lower().endswith('.zip') else MAX_FILE
    if len(content) > limit:
        raise AppError(ErrorCode.VALIDATION_ERROR, "ZIP 上限 50 MiB，单表格上限 10 MiB", 422)
    if filename.lower().endswith(".zip"):
        return [
            {**_inspect(data, name), "archive_member": name}
            for name, data in expand_zip(content, allow_template_notes=True)
        ]
    return [{**_inspect(content, filename), "archive_member": None}]
