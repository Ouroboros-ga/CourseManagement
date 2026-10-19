"""An XLSX adapter; student creation and enrollment changes are deliberately absent."""

from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.core.exceptions import AppError, ErrorCode

MAX_BYTES = 2 * 1024 * 1024


def parse_member_workbook(content: bytes) -> list[dict]:
    book = None
    try:
        if len(content) > MAX_BYTES:
            raise ValueError("文件最多2MiB")
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 128 or sum(e.file_size for e in entries) > 8 * 1024 * 1024:
                raise ValueError("表格解压大小超限")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        if len(book.worksheets) != 1:
            raise ValueError("仅接受一个工作表，请拆分文件后上传")
        sheet = book.worksheets[0]
        if (sheet.max_column or 0) > 16:
            raise ValueError("列数超限，请使用姓名、学号模板")
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        aliases = {
            "学号": "student_no",
            "student_no": "student_no",
            "姓名": "name",
            "学生姓名": "name",
            "name": "name",
        }
        columns = {}
        for i, cell in enumerate(header):
            key = aliases.get(str(cell.value or "").strip().lower())
            if key:
                if key in columns:
                    raise ValueError("学号或姓名表头重复")
                columns[key] = i
        if set(columns) != {"student_no", "name"}:
            raise ValueError("第一行必须包含学号、姓名两列")
        result, seen = [], set()
        for number, cells in enumerate(iterator, 2):
            if number > 1000:
                raise ValueError("表格行数超限，请清除多余格式行")
            if not any(c.value is not None for c in cells):
                continue
            values = {key: cells[index] for key, index in columns.items()}
            if any(c.data_type == "f" for c in values.values()):
                raise ValueError(f"第{number}行不能包含公式")
            if not isinstance(values["student_no"].value, str):
                raise ValueError(f"第{number}行学号须为文本，避免前导零丢失")
            no = values["student_no"].value.strip()
            name = str(values["name"].value or "").strip()
            if not no or not name or len(no) > 64 or len(name) > 128:
                raise ValueError(f"第{number}行学号与姓名必填且不能过长")
            if no in seen:
                raise ValueError(f"第{number}行学号重复，请每人仅保留一行")
            seen.add(no)
            result.append({"student_no": no, "name": name, "row": number})
            if len(result) > 200:
                raise ValueError("每次最多200名成员，请分批导入")
        if not result:
            raise ValueError("表格没有成员，请从第二行填写")
        return result
    except AppError:
        raise
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else "文件不是有效xlsx，请使用标准模板"
        raise AppError(ErrorCode.VALIDATION_ERROR, message, 422) from exc
    finally:
        if book is not None:
            book.close()


def member_template() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "志愿者组成员"
    sheet.append(["学号", "姓名"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="31577D")
    sheet["A1"].comment = Comment(
        "先建立有效学籍，再导入志愿者；学号按文本填写并保留前导零，每次最多200人。", "系统"
    )
    sheet["B1"].comment = Comment("与已有学生档案姓名一致；不会创建学生或自动更名。", "系统")
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 20
    for row in sheet.iter_rows(min_row=2, max_row=201, max_col=2):
        for cell in row:
            cell.number_format = "@"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:B201"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def workbook_payload(content, semester_id, effective_from, effective_to, enable_qualification):
    rows = parse_member_workbook(content)
    return {
        "semester_id": str(semester_id),
        "effective_from": effective_from,
        "effective_to": effective_to,
        "student_nos": [r["student_no"] for r in rows],
        "expected_names": {r["student_no"]: r["name"] for r in rows},
        "enable_qualification": enable_qualification,
        "skip_existing": True,
    }, {r["student_no"]: r["row"] for r in rows}
