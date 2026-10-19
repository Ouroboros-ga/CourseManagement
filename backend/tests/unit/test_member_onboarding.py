from datetime import date
from io import BytesIO

import pytest
from app.core.exceptions import AppError, ConflictError, PermissionDeniedError
from app.modules.academic.enrollment_models import StudentEnrollment
from app.modules.academic.models import Student, VolunteerQualification
from app.modules.volunteer.member_batch import GroupMemberBatchService
from app.modules.volunteer.member_models import VolunteerGroupMember
from app.modules.volunteer.member_service import GroupMemberService
from app.modules.volunteer.models import VolunteerGroup
from openpyxl import Workbook
from sqlalchemy import func, select

from tests.unit.test_campus_configuration import campus as campus_fixture
from tests.unit.test_group_business import prepare


@pytest.fixture(name="campus")
def fixture():
    yield from campus_fixture.__wrapped__()


def batch(campus):
    _, db, admin, a, _, semester, group, _, student, _ = prepare(campus)
    return (
        db,
        student,
        GroupMemberBatchService(db, admin, int(a["id"]), int(group["id"])),
        {
            "semester_id": str(semester.id),
            "student_nos": [student.student_no],
            "expected_names": {student.student_no: student.name},
            "enable_qualification": True,
            "skip_existing": True,
            "effective_from": "2026-09-07",
            "effective_to": "2027-01-01",
        },
    )


def test_preview_does_not_write_and_confirm_enables_qualification_with_membership(campus):
    db, student, service, body = batch(campus)
    preview = service.preview(body)
    assert preview["can_commit"] and preview["qualifications_to_enable"] == 1
    assert db.scalar(select(func.count()).select_from(VolunteerQualification)) == 0
    assert db.scalar(select(func.count()).select_from(VolunteerGroupMember)) == 0
    result = service.commit(body, preview["revision"])
    assert result["created"] == 1 and result["qualifications_enabled"] == 1
    qual = db.scalar(
        select(VolunteerQualification).where(VolunteerQualification.student_id == student.id)
    )
    assert qual.enabled and qual.college_id == service.cid
    repeated = service.commit(body, service.preview(body)["revision"])
    assert repeated["created"] == 0 and repeated["unchanged"] == 1


def test_names_and_unknown_numbers_block_whole_batch_without_creating_students(campus):
    db, student, service, body = batch(campus)
    count = db.scalar(select(func.count()).select_from(Student))
    wrong = {
        **body,
        "student_nos": [student.student_no, "00000000000"],
        "expected_names": {student.student_no: "Wrong name", "00000000000": "Missing"},
    }
    preview = service.preview(wrong)
    assert not preview["can_commit"] and len(preview["issues"]) == 2
    with pytest.raises(ConflictError):
        service.commit(wrong, preview["revision"])
    assert db.scalar(select(func.count()).select_from(Student)) == count
    assert db.scalar(select(func.count()).select_from(VolunteerGroupMember)) == 0
    assert db.scalar(select(func.count()).select_from(VolunteerQualification)) == 0


def test_non_enrolled_student_cannot_be_imported(campus):
    db, student, service, body = batch(campus)
    db.add(
        StudentEnrollment(
            student_id=student.id,
            college_id=service.cid,
            administrative_class_id=student.administrative_class_id,
            effective_from=date(2026, 9, 1),
            status="SUSPENDED",
            revision=1,
            snapshot={},
        )
    )
    db.commit()
    assert not service.preview(body)["can_commit"]


def test_inactive_group_can_prepare_members_but_stale_preview_cannot_commit(campus):
    db, student, service, body = batch(campus)
    group = db.get(VolunteerGroup, service.gid)
    group.status = "INACTIVE"
    db.commit()
    preview = service.preview(body)
    assert preview["can_commit"]
    student.name = "Changed"
    student.lock_version += 1
    db.commit()
    with pytest.raises(ConflictError):
        service.commit(body, preview["revision"])
    assert db.scalar(select(func.count()).select_from(VolunteerGroupMember)) == 0


def test_single_add_to_group_can_atomically_enable_qualification(campus):
    _, db, admin, a, _, semester, group, _, student, _ = prepare(campus)
    result = GroupMemberService(db, admin, int(a["id"]), int(group["id"])).create(
        {
            "student_id": str(student.id),
            "semester_id": str(semester.id),
            "effective_from": "2026-09-07",
            "effective_to": "2027-01-01",
            "enable_qualification": True,
        }
    )
    assert result["student_id"] == str(student.id)
    assert db.scalar(select(VolunteerQualification.enabled)) is True


def test_membership_failure_rolls_back_qualification_and_account(campus, monkeypatch):
    from app.modules.identity.models import UserAccount

    db, _, service, body = batch(campus)
    accounts = db.scalar(select(func.count()).select_from(UserAccount))
    preview = service.preview(body)

    def fail(*args, **kwargs):
        raise ConflictError(message="simulated member write failure")

    monkeypatch.setattr(service.members, "_save", fail)
    with pytest.raises(ConflictError):
        service.commit(body, preview["revision"])
    assert db.scalar(select(func.count()).select_from(VolunteerQualification)) == 0
    assert db.scalar(select(func.count()).select_from(VolunteerGroupMember)) == 0
    assert db.scalar(select(func.count()).select_from(UserAccount)) == accounts


def test_one_student_can_join_two_groups_without_duplicating_identity(campus):
    _, db, actor, a, _, sem, one, two, student, _ = prepare(campus)
    for group in (one, two):
        GroupMemberService(db, actor, int(a["id"]), int(group["id"])).create(
            {
                "student_id": str(student.id),
                "semester_id": str(sem.id),
                "effective_from": "2026-09-07",
                "effective_to": "2027-01-01",
                "enable_qualification": True,
            }
        )
    assert db.scalar(select(func.count()).select_from(VolunteerGroupMember)) == 2
    assert db.scalar(select(func.count()).select_from(VolunteerQualification)) == 1


def test_onboarding_never_attaches_an_unrelated_account_by_username(campus):
    from app.modules.identity.models import UserAccount

    db, student, service, body = batch(campus)
    other = UserAccount(
        username=student.student_no, display_name="Unrelated account", status="ACTIVE"
    )
    db.add(other)
    db.commit()
    service.commit(body, service.preview(body)["revision"])
    assert other.student_id is None
    assert db.scalar(select(UserAccount.id).where(UserAccount.student_id == student.id)) != other.id


def test_member_manager_cannot_grant_qualification_without_its_permission(campus):
    from tests.unit.test_campus_configuration import grant

    db, student, service, body = batch(campus)
    config, _, admin, reader, a, _ = campus
    grant(config, admin, reader, a, ["volunteer.member.manage"])
    restricted = GroupMemberBatchService(db, reader, service.cid, service.gid)
    with pytest.raises(PermissionDeniedError):
        restricted.commit(body, "0" * 64)
    assert db.scalar(select(func.count()).select_from(VolunteerQualification)) == 0


def workbook(rows):
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_excel_preserves_text_ids_and_requires_names():
    from app.modules.volunteer.member_excel import parse_member_workbook

    assert parse_member_workbook(workbook([["学号", "姓名"], ["00123456", "张三"]])) == [
        {"student_no": "00123456", "name": "张三", "row": 2}
    ]
    for rows in (
        [["学号", "姓名"], [123456, "张三"]],
        [["学号"], ["00123456"]],
        [["学号", "姓名"], ["00123456", "张三"], ["00123456", "张三"]],
    ):
        with pytest.raises(AppError):
            parse_member_workbook(workbook(rows))


def test_excel_rejects_formula_and_multiple_sheets():
    from app.modules.volunteer.member_excel import parse_member_workbook

    with pytest.raises(AppError):
        parse_member_workbook(workbook([["学号", "姓名"], ["00123456", "=1+1"]]))
    book = Workbook()
    book.active.append(["学号", "姓名"])
    book.create_sheet("Other")
    output = BytesIO()
    book.save(output)
    with pytest.raises(AppError):
        parse_member_workbook(output.getvalue())
