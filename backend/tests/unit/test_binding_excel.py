from io import BytesIO

from openpyxl import load_workbook


def test_shared_binding_workbook_preserves_text_and_student_layout():
    from app.modules.identity.binding_excel import build_binding_workbook

    content = build_binding_workbook(
        title="学生绑定码",
        headers=["学号", "姓名", "绑定码"],
        rows=[["00001", "=测试", "012ABC"]],
        code_column=3,
        widths=[16, 12, 14],
    )
    sheet = load_workbook(BytesIO(content)).active
    assert sheet.title == "学生绑定码"
    assert [cell.value for cell in sheet[2]] == ["00001", "=测试", "012ABC"]
    assert all(cell.data_type == "s" for cell in sheet[2])
    assert sheet["A1"].fill.fgColor.rgb == "001890FF"


def test_student_export_keeps_columns_and_beijing_expiry(monkeypatch):
    from app.modules.identity.service import IdentityService

    monkeypatch.setattr(
        IdentityService,
        "batch_issue_binding_tokens",
        lambda *args, **kwargs: [
            {
                "student_no": "00001",
                "name": "测试学生",
                "class_name": "测试班",
                "college": "测试学院",
                "binding_code": "012ABC",
                "expires_at": "2026-10-21T10:00:00Z",
                "is_bound": False,
            }
        ],
    )
    # 渲染用例不访问数据库；签发授权由原有业务用例覆盖。
    content = IdentityService.__new__(IdentityService).export_binding_tokens_excel(actor_user_id=1)
    sheet = load_workbook(content).active
    assert sheet.title == "学生绑定码"
    assert [cell.value for cell in sheet[2]] == [
        "00001",
        "测试学生",
        "测试班",
        "测试学院",
        "012ABC",
        "2026-10-21 18:00:00",
        "未绑定",
    ]
