"""查询优化必须同时保持授权结果与受控的 SQL/内存成本。只使用内存库。"""

from contextlib import contextmanager
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from app.common.pagination import PageParams
from app.modules.academic.staff_models import TeachingStaffAssignment
from app.modules.attendance.models import AttendanceRecord
from app.modules.inspection.models import InspectionTask, TaskRosterMember
from app.modules.inspection.staff_record_query import StaffRecordQueryService
from sqlalchemy import event

from tests.unit import test_occurrence_location as loc
from tests.unit import test_staff_corrections as corr
from tests.unit.test_staff_records import _teacher_actor
from tests.unit.test_staff_records import staff_records_db as records_fixture


@pytest.fixture
def staff_records_db():
    yield from records_fixture.__wrapped__()


@contextmanager
def statements(db):
    queries = []

    def collect(_conn, _cursor, statement, _parameters, _context, _many):
        if statement.lstrip().upper().startswith("SELECT"):
            queries.append(statement)

    engine = db.get_bind()
    event.listen(engine, "before_cursor_execute", collect)
    try:
        yield queries
    finally:
        event.remove(engine, "before_cursor_execute", collect)


@pytest.fixture
def query(staff_records_db):
    db = staff_records_db
    db.add(TeachingStaffAssignment(
        staff_id=101, semester_id=1, anchor_type="LOCAL", teaching_class_id=1,
        effective_from=date(2026, 9, 1),
    ))
    for i in range(10, 30):
        db.add(InspectionTask(
            id=i, task_key=f"BATCH-{i}", college_id=1, semester_id=1,
            inspection_date=date(2026, 10, 18), week_no=8,
            inspection_type="COURSE", start_period=1, end_period=2,
            teaching_class_id=1, roster_version=1,
        ))
    db.flush()
    for i in range(10, 30):
        db.add(TaskRosterMember(
            task_id=i, roster_version=1, student_id=1, student_no="S001", name="张三",
            administrative_class_id_snapshot=1,
        ))
        db.add(AttendanceRecord(
            task_id=i, student_id=1, effective_type="NORMAL", current_version=1,
        ))
    db.commit()
    return db, StaffRecordQueryService(db, _teacher_actor())


def test_card_query_count_does_not_grow_with_page_size(query):
    db, svc = query
    with statements(db) as small:
        svc.list_cards(
            "teaching", PageParams(page=1, page_size=1), inspection_date=date(2026, 10, 18)
        )
    with statements(db) as large:
        cards, total = svc.list_cards(
            "teaching", PageParams(page=1, page_size=50), inspection_date=date(2026, 10, 18)
        )
    assert total == 20
    assert sum(c.counts.total for c in cards if c.counts) == 20
    assert len(large) <= len(small) + 2, (len(small), len(large))


def test_teaching_summary_aggregates_counts_in_sql(query):
    db, svc = query
    with statements(db) as queries:
        summary = svc.get_summary("teaching")
    assert summary.student_counts.total == 23
    assert summary.recorded_task_count == 21
    assert any("GROUP BY" in sql and "attendance_record.effective_type" in sql for sql in queries)


def test_export_streams_only_the_current_roster_version(query, monkeypatch):
    db, svc = query
    db.add(TaskRosterMember(
        task_id=1, roster_version=0, student_id=1, student_no="OLD-NUMBER", name="旧姓名",
    ))
    db.commit()
    from app.modules.inspection import staff_record_query as module

    original = module.build_staff_export_excel

    def capture(rows):
        assert not isinstance(rows, list), "Export must not accumulate all detail rows"
        return original(rows)

    monkeypatch.setattr(module, "build_staff_export_excel", capture)
    with statements(db) as queries:
        data, _ = svc.export_excel("teaching")
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
    values = list(wb.active.values)
    wb.close()
    assert len(values) == 24
    assert all(row[8] != "OLD-NUMBER" for row in values)
    assert any(
        "task_roster_member.roster_version = inspection_task.roster_version" in sql
        for sql in queries
    )


def test_calendar_avoids_per_task_queries_for_unrelated_history():
    with contextmanager(loc.location_db.__wrapped__)() as db:
        loc.add_relation(db)
        own = loc.make_task(db, key="OWN", with_source=True)
        for i in range(15):
            other = loc.make_task(db, key=f"OTHER-{i}")
            other.teaching_class_id = 2
            other.course_schedule_id = 2
        db.commit()
        actor = loc.staff_ref(db)
        with statements(db) as queries:
            result = loc.service(db).occurrence_calendar(
                actor, date_from=loc.COURSE_DAY, date_to=loc.COURSE_DAY,
            )
        assert result.total == 1
        assert result.items[0].task_ids == [str(own.id)]
        source_queries = [s for s in queries if "FROM task_schedule_source" in s]
        assert len(source_queries) <= 2, len(source_queries)


@pytest.mark.parametrize("provenance,shared", [(True, False), (False, False), (True, True)])
def test_sql_aggregates_preserve_merged_class_visibility(provenance, shared):
    with contextmanager(corr.correction_db.__wrapped__)() as db:
        corr.add_relation(db, local_class=None if shared else 1, shared=shared)
        corr.make_merged_task(db, provenance=provenance)
        task = db.get(InspectionTask, 3)
        task.course_name_snapshot = "ONLY-MERGED"
        db.commit()
        svc, _ = corr._query(db)
        detail, total = svc.list_attendance(3, "teaching", None, PageParams(page=1, page_size=50))
        assert svc.get_card(3, "teaching").counts.total == total
        assert svc.get_summary("teaching", course_name="ONLY-MERGED").student_counts.total == total
        data, _ = svc.export_excel("teaching", course_name="ONLY-MERGED")
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
        exported = list(wb.active.values)[1:]
        wb.close()
        assert {row[8] for row in exported} == {row.student_no for row in detail}


def test_calendar_matches_shared_snapshot_without_a_local_source():
    with contextmanager(loc.location_db.__wrapped__)() as db:
        loc.add_relation(db)
        task = loc.make_task(db, key="SHARED-PEER", snapshot=[loc.KEY])
        task.teaching_class_id = 2
        task.course_schedule_id = 2
        db.commit()
        page = loc.calendar(db, loc.COURSE_DAY, loc.COURSE_DAY)
        assert page.items[0].task_ids == [str(task.id)]


def test_calendar_numeric_shared_key_is_not_coerced_to_a_bigint():
    from app.modules.academic.models import CourseSchedule
    from sqlalchemy import select

    with contextmanager(loc.location_db.__wrapped__)() as db:
        loc.add_relation(db)
        meeting = "s:" + "9" * 30
        for schedule in db.scalars(select(CourseSchedule)):
            schedule.shared_meeting_key = meeting
        task = loc.make_task(db, key="LONG-KEY", snapshot=[f"{meeting}|{loc.COURSE_DAY}"])
        page = loc.calendar(db, loc.COURSE_DAY, loc.COURSE_DAY)
        assert page.items[0].task_ids == [str(task.id)]


def test_write_only_export_preserves_chinese_styles_and_leading_zero():
    from app.modules.report.staff_export import build_staff_export_excel

    data = build_staff_export_excel(iter([{
        "student_no": "001234", "student_name": "测试姓名", "current_version": 7,
    }]))
    wb = openpyxl.load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.title == "查课考勤记录"
    assert ws["I2"].value == "001234" and ws["I2"].data_type == "s"
    assert ws["M2"].value == 7
    assert ws["A1"].fill.fgColor.rgb == "001F4E79"
    assert ws["J2"].font.name == "Microsoft YaHei"
    assert ws.column_dimensions["J"].width >= 12
    wb.close()
