from hashlib import sha256
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import pytest
from app.core.config import get_settings
from app.core.exceptions import AppError
from app.modules.agent.gateway import invoke_tool
from app.modules.agent.run_context import AppConnection, RunService
from app.modules.inspection.models import InspectionTask
from sqlalchemy import select
from tests.unit.agent.test_run_context import session_for
from tests.unit.test_campus_configuration import campus as campus_fixture
from tests.unit.test_scoped_file_access import setup


@pytest.fixture(name="case")
def fixture(monkeypatch):
    original = campus_fixture.__wrapped__()
    campus = next(original)
    monkeypatch.setattr(get_settings(), "campus_configuration_enabled", True)
    actor, _ = setup(campus, monkeypatch)
    db = campus[1]
    auth = session_for(db, actor)
    app = AppConnection(
        app_key_id="management",
        entry_kind="MANAGEMENT",
        environment="dev",
        key_digest=sha256(("m" * 40).encode()).hexdigest(),
    )
    runs = RunService(db, [app], environment="dev")
    task = db.scalar(
        select(InspectionTask).where(InspectionTask.college_id == int(campus[4]["id"]))
    )
    _, token = runs.issue(
        "management",
        actor.id,
        auth.id,
        college_id=int(campus[4]["id"]),
        semester_id=task.semester_id,
    )
    yield campus, actor, runs, token, task
    with pytest.raises(StopIteration):
        next(original)


def prepare(case, **extra):
    campus, _, runs, token, task = case
    return invoke_tool(
        campus[1],
        runs.resolve("m" * 40, token),
        "report_attendance_export_prepare",
        {"semester_id": str(task.semester_id), "week_no": task.week_no, **extra},
    ).data


@pytest.mark.parametrize(
    "format_name,suffix",
    [
        ("ANALYTICS_XLSX", "/export"),
        ("DETAIL_DOCX", "/export-detail-doc"),
        ("STUDY_XLSX", "/export-study-excel"),
    ],
)
def test_preparation_is_a_scoped_authenticated_action_not_a_generated_artifact(
    case, format_name, suffix
):
    data = prepare(case, format=format_name)
    assert data["artifact_status"] == "NOT_GENERATED"
    assert data["download"]["requires_login"] is True
    url = urlsplit(data["download"]["path"])
    assert not url.scheme and not url.netloc
    assert url.path == f"/api/v1/colleges/{case[0][4]['id']}/statistics/attendance{suffix}"
    assert parse_qs(url.query) == {
        "semester_id": [str(case[4].semester_id)],
        "week_no": [str(case[4].week_no)],
    }
    assert "sig=" not in str(data) and "token" not in str(data)
    assert data["matching_task_count"] == 1


@pytest.mark.parametrize(
    "extra",
    [
        {"college_id": "999"},
        {"student_id": "999"},
        {"url": "https://untrusted.invalid"},
        {"format": "SQL"},
        {"semester_id": "01"},
    ],
)
def test_model_cannot_choose_identity_or_arbitrary_destinations(case, extra):
    with pytest.raises(AppError) as error:
        prepare(case, **extra)
    assert error.value.http_status == 422


@pytest.mark.parametrize("format_name", ["ANALYTICS_XLSX", "DETAIL_DOCX", "STUDY_XLSX"])
def test_revoke_after_preparation_blocks_actual_download(case, monkeypatch, format_name):
    from app.core import database
    from app.main import create_app
    from app.modules.academic.enrollment_models import TeacherCollegeBinding
    from app.modules.identity.deps import get_current_user
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from sqlalchemy.orm import sessionmaker

    campus, actor, _, _, _ = case
    own = case[4]
    own.course_name_snapshot = "OWN-EXPORT-MARKER"
    foreign = campus[1].scalar(
        select(InspectionTask).where(InspectionTask.college_id == int(campus[5]["id"]))
    )
    foreign.course_name_snapshot = "FOREIGN-EXPORT-MARKER"
    campus[1].commit()
    data = prepare(case, format=format_name)
    db = campus[1]
    factory = sessionmaker(bind=db.bind, expire_on_commit=False)
    monkeypatch.setattr(database, "get_session_factory", lambda: factory)
    app = create_app()
    app.dependency_overrides[get_current_user] = lambda: actor
    client = TestClient(app)
    response = client.get(data["download"]["path"])
    assert response.status_code == 200
    if format_name == "DETAIL_DOCX":
        from docx import Document

        document = Document(BytesIO(response.content))
        values = [p.text for p in document.paragraphs]
        values += [
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        ]
        assert document.paragraphs or document.tables
    else:
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        assert workbook.sheetnames
        values = [
            str(value)
            for sheet in workbook
            for row in sheet.iter_rows(values_only=True)
            for value in row
            if value is not None
        ]
        workbook.close()
    if format_name == "ANALYTICS_XLSX":
        assert "OWN-EXPORT-MARKER" in values
    assert "FOREIGN-EXPORT-MARKER" not in " ".join(values)
    db.delete(db.get(TeacherCollegeBinding, actor.id))
    db.commit()
    assert client.get(data["download"]["path"]).status_code in {403, 404}


def test_queue_visibility_uses_statistics_scope_instead_of_a_student_binding(case):
    from app.modules.agent.task_service import visibility_hash

    campus, _, runs, token, _ = case
    context = runs.resolve("m" * 40, token)
    assert visibility_hash(
        campus[1], context, "report_attendance_export_prepare"
    ) == visibility_hash(campus[1], context, "report_attendance_summary")


def test_business_gate_is_checked_before_a_download_action_is_issued(case, monkeypatch):
    monkeypatch.setattr(get_settings(), "campus_business_scoped_enabled", False)
    with pytest.raises(AppError):
        prepare(case)


def test_date_range_action_is_bounded_and_does_not_include_week_number(case):
    day = case[4].inspection_date.isoformat()
    data = prepare(case, week_no=None, date_from=day, date_to=day)
    assert parse_qs(urlsplit(data["download"]["path"]).query) == {
        "semester_id": [str(case[4].semester_id)],
        "date_from": [day],
        "date_to": [day],
    }


def test_run_bound_semester_cannot_be_replaced_by_the_model(case):
    with pytest.raises(AppError) as error:
        prepare(case, semester_id=str(case[4].semester_id + 1))
    assert error.value.http_status == 403


def test_real_queue_worker_can_prepare_export_for_a_teacher_without_student_binding(case):
    from app.modules.agent.models import AgentRun
    from app.modules.agent.task_service import TaskService
    from app.modules.agent.worker import Worker
    from sqlalchemy.orm import sessionmaker

    campus, actor, runs, token, source = case
    db = campus[1]
    context = runs.resolve("m" * 40, token)
    run = db.get(AgentRun, int(context.run_id))
    service = TaskService(db, runs.connections, environment="dev")
    task = service.submit(
        run.id,
        actor.id,
        run.source_session_id,
        "report_attendance_export_prepare",
        {"semester_id": str(source.semester_id), "week_no": source.week_no},
        "export-1",
    )
    db.commit()
    assert Worker(sessionmaker(bind=db.bind), runs.connections, environment="dev").run_once()
    state = service.read(task.id, actor.id)
    assert state["status"] == "SUCCEEDED"
    assert state["result"]["data"]["artifact_status"] == "NOT_GENERATED"


def test_tool_http_uses_actual_run_credentials_and_returns_metadata_only(case, monkeypatch):
    from app.core import database
    from app.main import create_app
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import sessionmaker

    campus, _, runs, token, source = case
    settings = get_settings()
    monkeypatch.setattr(settings, "agent_tools_enabled", True)
    monkeypatch.setattr(settings, "agent_tool_connections", runs.connections)
    factory = sessionmaker(bind=campus[1].bind, expire_on_commit=False)
    monkeypatch.setattr(database, "get_session_factory", lambda: factory)
    client = TestClient(create_app())
    path = "/api/v1/agent-tools/v1/calls/report_attendance_export_prepare"
    body = {"arguments": {"semester_id": str(source.semester_id), "week_no": source.week_no}}
    assert client.post(path, json=body).status_code == 401
    response = client.post(
        path,
        json=body,
        headers={
            "X-Tool-App-Key": "m" * 40,
            "X-Agent-Run-Token": token,
        },
    )
    assert response.status_code == 200
    assert response.json()["data"]["data"]["artifact_status"] == "NOT_GENERATED"
    assert "FOREIGN" not in response.text and "file_id" not in response.text


def test_self_service_directory_does_not_acquire_report_export():
    from app.modules.agent.registry import ToolRegistry

    assert "report_attendance_export_prepare" not in {
        item.name for item in ToolRegistry().list_for("SELF_SERVICE")
    }
    assert "report_attendance_export_prepare" not in {
        item.name for item in ToolRegistry().list_for("AUTOMATION", "objection")
    }


def test_preparation_does_not_publish_report_versions_or_files(case):
    from app.modules.file.models import FileObject
    from app.modules.report.models import ReportVersion

    db = case[0][1]
    before = db.query(FileObject).count(), db.query(ReportVersion).count()
    prepare(case)
    assert (db.query(FileObject).count(), db.query(ReportVersion).count()) == before
