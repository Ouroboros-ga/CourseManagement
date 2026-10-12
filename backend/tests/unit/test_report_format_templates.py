"""验证导出 Word 和 Excel 表格格式 100% 对齐用户模板。

对比参照源：
- media_1791768520700_0eb30860.docx（课堂考勤详表：居中标题'详表'，5列[日期, 班级, 课程, 姓名/学号, 迟/退/旷]）
- media_1791768559858_5aaaadb8.xlsx（自习与日常考勤：4列[班级, 姓名, 情况, 日期]）
"""

import io
import docx
import openpyxl
from docx.shared import Pt

from app.modules.report.artifacts import (
    build_detail_table_word_document,
    build_study_attendance_excel,
    build_word_document,
    build_excel,
    build_attendance_analytics_excel,
)


def _sample_payload():
    return {
        "meta": {"semester_id": 1, "week_no": 2, "scope": "COLLEGE", "eligible_task_count": 8},
        "overall": {
            "statistics_available": True,
            "expected_count": 120,
            "normal_count": 110,
            "abnormal_count": 10,
            "leave_count": 2,
            "late_count": 3,
            "absent_count": 5,
            "abnormal_rate": 0.0833,
        },
        "classes": [
            {
                "class_name": "机电2401",
                "expected_count": 60,
                "normal_count": 55,
                "late_count": 2,
                "absent_count": 3,
                "leave_count": 0,
                "abnormal_rate": 0.0833,
                "class_ratio_available": True,
            },
            {
                "class_name": "智能制造2501",
                "expected_count": 60,
                "normal_count": 55,
                "late_count": 1,
                "absent_count": 2,
                "leave_count": 2,
                "abnormal_rate": 0.0833,
                "class_ratio_available": True,
            },
        ],
        "tasks": [
            {
                "task_id": "101",
                "inspection_date": "2026-03-16",
                "start_period": 1,
                "end_period": 2,
                "course_name_snapshot": "单片机原理及应用技术",
                "class_name_snapshot": "机电2401",
                "classroom_snapshot": "实训楼401",
                "expected_count": 60,
                "normal_count": 55,
                "late_count": 2,
                "absent_count": 3,
                "leave_count": 0,
                "abnormal_count": 5,
                "attendance_rate": 0.9167,
            }
        ],
        "details": [
            {
                "inspection_date": "2026-03-16",
                "date_display": "3.16",
                "class_name_snapshot": "机电2401",
                "course_name": "单片机原理及应用技术",
                "name": "王树涵",
                "student_no": "20240101",
                "attendance_type": "ABSENT",
                "inspection_type": "COURSE",
            },
            {
                "inspection_date": "2026-03-16",
                "date_display": "3.16",
                "class_name_snapshot": "机电2401",
                "course_name": "单片机原理及应用技术",
                "name": "梁敏琪",
                "student_no": "20240102",
                "attendance_type": "LATE",
                "inspection_type": "COURSE",
            },
            {
                "inspection_date": "2026-03-18",
                "date_display": "3.18",
                "class_name_snapshot": "智能制造2501",
                "course_name": "",
                "name": "周家乐",
                "student_no": "20250101",
                "attendance_type": "ABSENT",
                "inspection_type": "EVENING_STUDY",
            },
        ],
    }


def test_build_detail_table_word_document():
    payload = _sample_payload()
    data = build_detail_table_word_document(payload)
    doc = docx.Document(io.BytesIO(data))

    # 1. 验证主标题为居中「详表」
    p_titles = [p for p in doc.paragraphs if p.text.strip()]
    assert len(p_titles) >= 1
    assert "详表" in p_titles[0].text

    # 2. 验证表格结构
    assert len(doc.tables) >= 1
    tbl = doc.tables[0]
    headers = [c.text.strip() for c in tbl.rows[0].cells]
    assert headers == ["日期", "班级", "课程", "姓名/学号", "迟/退/旷"]

    # 验证课堂考勤明细数据与考勤缩写
    r1 = [c.text.strip() for c in tbl.rows[1].cells]
    assert r1[0] == "3.16"
    assert r1[1] == "机电2401"
    assert r1[2] == "单片机原理及应用技术"
    assert r1[3] == "王树涵"
    assert r1[4] == "旷"  # 短标签对齐

    # 验证列宽精确对应
    expected_widths = [Pt(79), Pt(94), Pt(162), Pt(78), Pt(63)]
    for idx, exp_w in enumerate(expected_widths):
        # 允许 1pt 内的微小浮点误差
        assert abs(tbl.rows[0].cells[idx].width.pt - exp_w.pt) < 1.0

    # 3. 验证自习考勤表结构（存在晚自习记录时）
    assert len(doc.tables) == 2
    std_tbl = doc.tables[1]
    std_headers = [c.text.strip() for c in std_tbl.rows[0].cells]
    assert std_headers == ["班级", "姓名", "情况", "日期"]
    std_r1 = [c.text.strip() for c in std_tbl.rows[1].cells]
    assert std_r1 == ["智能制造2501", "周家乐", "旷课", "3.18"]


def test_build_study_attendance_excel():
    payload = _sample_payload()
    data = build_study_attendance_excel(payload)
    wb = openpyxl.load_workbook(io.BytesIO(data))

    assert "Sheet1" in wb.sheetnames
    ws = wb["Sheet1"]
    headers = [ws.cell(1, c).value for c in range(1, 5)]
    assert headers == ["班级", "姓名", "情况", "日期"]

    # 验证列宽对齐原版 xlsx
    assert ws.column_dimensions["A"].width == 14.11
    assert ws.column_dimensions["B"].width == 13.0
    assert ws.column_dimensions["C"].width == 13.0
    assert ws.column_dimensions["D"].width == 13.0

    # 验证行内容
    row2 = [ws.cell(2, c).value for c in range(1, 5)]
    assert row2 == ["智能制造2501", "周家乐", "旷课", "3.18"]


def test_build_comprehensive_excel_and_word():
    payload = _sample_payload()

    # 验证周报 Word 包含详表与自习表
    word_data = build_word_document(payload)
    doc = docx.Document(io.BytesIO(word_data))
    assert len(doc.tables) >= 3  # KPI表、班级表、详表、自习表
    det_tbl = doc.tables[2]
    headers = [c.text.strip() for c in det_tbl.rows[0].cells]
    assert headers == ["日期", "班级", "课程", "姓名/学号", "迟/退/旷"]

    # 验证周报 Excel
    weekly_excel = build_excel(payload)
    wb_w = openpyxl.load_workbook(io.BytesIO(weekly_excel))
    assert "详表" in wb_w.sheetnames
    assert "自习与日常考勤" in wb_w.sheetnames
    assert [wb_w["详表"].cell(1, c).value for c in range(1, 6)] == ["日期", "班级", "课程", "姓名/学号", "迟/退/旷"]
    assert [wb_w["自习与日常考勤"].cell(1, c).value for c in range(1, 5)] == ["班级", "姓名", "情况", "日期"]

    # 验证多维分析看板 Excel
    analytics_excel = build_attendance_analytics_excel(payload)
    wb_a = openpyxl.load_workbook(io.BytesIO(analytics_excel))
    assert "学院考勤总览与看板" in wb_a.sheetnames
    assert "行政班到课率统计" in wb_a.sheetnames
    assert "教学班与课堂到课率明细" in wb_a.sheetnames
    assert "详表" in wb_a.sheetnames
    assert "自习与日常考勤" in wb_a.sheetnames
