"""课表导出产物必须保留范围和 Excel 文本安全性。"""

from io import BytesIO

import pytest
from app.modules.academic.export_service import build_workbook
from openpyxl import load_workbook


def test_export_columns_and_formula_like_text_are_literal() -> None:
    rows = [("=某教学班", "+课程", "-教师", "1,3", 2, 1, 2, "@A101")]
    book = load_workbook(BytesIO(build_workbook(rows)), read_only=True)
    sheet = book.active
    assert sheet is not None
    assert list(next(sheet.values)) == [
        "教学班", "课程", "教师", "周次", "星期", "开始节次", "结束节次", "教室"
    ]
    cells = list(sheet.iter_rows())[1]
    assert [cell.value for cell in cells] == list(rows[0])
    for cell in (cells[0], cells[1], cells[2], cells[3], cells[7]):
        assert cell.data_type == "s"


def test_export_rejects_more_than_ten_thousand_rows() -> None:
    with pytest.raises(ValueError, match="10000"):
        build_workbook([("班", "课", "", "1", 1, 1, 2, "") for _ in range(10001)])


def test_export_empty_result_still_has_header() -> None:
    book = load_workbook(BytesIO(build_workbook([])), read_only=True)
    sheet = book.active
    assert sheet is not None
    assert len(list(sheet.values)) == 1
