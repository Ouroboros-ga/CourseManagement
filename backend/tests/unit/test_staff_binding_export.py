"""绑定码 Excel 导出使用独立内存库；不签发真实教师凭证。"""

from datetime import timedelta
from io import BytesIO

import pytest
from app.core.database import utcnow
from app.core.exceptions import ConflictError, PermissionDeniedError
from app.core.security import hash_token
from app.modules.academic.staff_models import FacultyStaff
from app.modules.audit.models import AuditLog
from app.modules.identity.staff_binding import StaffBindingService
from app.modules.identity.staff_binding_models import StaffBindingToken
from openpyxl import load_workbook
from sqlalchemy import select

from tests.unit.test_staff_management import create_staff
from tests.unit.test_staff_management import staff_env as staff_env_fixture


@pytest.fixture
def staff_env():
    yield from staff_env_fixture.__wrapped__()


def export(env, **kwargs):
    return StaffBindingService(env["db"], college_id=1).export_tokens_excel(
        actor=env["admin"], **kwargs
    )


def test_export_filters_scope_and_preserves_text(staff_env):
    row = create_staff(staff_env, "00023", "导出教师")
    create_staff(staff_env, "00024", "其他教师")
    create_staff(staff_env, "00025", "导出外院", college_id=2)
    db = staff_env["db"]
    blob = export(staff_env, keyword="导出")
    sheet = load_workbook(BytesIO(blob)).active
    assert sheet.max_row == 2
    assert [cell.value for cell in sheet[1]][:4] == ["工号", "姓名", "所属学院", "一次性绑定码"]
    assert sheet["A2"].value == "00023" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == "导出教师" and sheet["C2"].value == "甲院"
    code = sheet["D2"].value
    assert isinstance(code, str) and len(code) == 6 and sheet["D2"].data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:E2"
    token = db.scalar(select(StaffBindingToken))
    assert token.staff_id == int(row["id"]) and token.token_hash == hash_token(code)
    assert token.status == "UNUSED"
    assert utcnow() + timedelta(days=29) < token.expires_at < utcnow() + timedelta(days=31)
    audits = db.scalars(
        select(AuditLog).where(AuditLog.action == "staff_binding_token.issue")
    ).all()
    assert len(audits) == 1 and code not in str(audits[0].after_json)


@pytest.mark.parametrize("invalid", ["foreign", "inactive", "linked"])
def test_selected_ineligible_rejects_whole_batch(staff_env, invalid):
    valid = create_staff(staff_env, "00001")
    other = create_staff(staff_env, "00002", college_id=2 if invalid == "foreign" else 1)
    db = staff_env["db"]
    profile = db.get(FacultyStaff, int(other["id"]))
    if invalid == "inactive":
        profile.status = "INACTIVE"
    elif invalid == "linked":
        profile.user_id = staff_env["accounts"]["wechat"].id
    db.commit()
    with pytest.raises(ConflictError):
        export(staff_env, staff_ids=[int(valid["id"]), int(other["id"])])
    assert db.scalar(select(StaffBindingToken.id)) is None


def test_export_render_failure_rolls_back_issuance(staff_env, monkeypatch):
    create_staff(staff_env, "00001")

    def fail(*args, **kwargs):
        raise OSError("模拟文件生成失败")

    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", fail)
    with pytest.raises(OSError):
        export(staff_env)
    db = staff_env["db"]
    assert db.scalar(select(StaffBindingToken.id)) is None
    assert (
        db.scalar(select(AuditLog.id).where(AuditLog.action == "staff_binding_token.issue")) is None
    )


def test_export_requires_authority_and_has_bounded_batch(staff_env):
    service = StaffBindingService(staff_env["db"], college_id=1)
    with pytest.raises(PermissionDeniedError):
        service.export_tokens_excel(actor=staff_env["plain"])
    with pytest.raises(ConflictError):
        export(staff_env)
    with pytest.raises(ConflictError):
        export(staff_env, staff_ids=list(range(1, 502)))


def test_exported_code_redeems_through_existing_staff_binding(staff_env):
    row = create_staff(staff_env, "00001")
    sheet = load_workbook(
        BytesIO(export(staff_env, staff_ids=[int(row["id"])], days_valid=14))
    ).active
    code = sheet["D2"].value
    binding = StaffBindingService(staff_env["db"], college_id=1)
    binding.redeem(user_id=staff_env["plain"].id, staff_no="00001", binding_code=code)
    assert staff_env["db"].get(FacultyStaff, int(row["id"])).user_id == staff_env["plain"].id
    with pytest.raises(ConflictError):
        binding.redeem(user_id=staff_env["plain"].id, staff_no="00001", binding_code=code)


def test_download_route_and_request_validation(staff_env):
    from app.core.database import get_db
    from app.modules.academic import staff_router
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    create_staff(staff_env, "00001")
    app = FastAPI()
    app.include_router(staff_router.router, prefix="/api/v1")
    app.dependency_overrides[staff_router.require_campus_enabled] = lambda: None
    app.dependency_overrides[staff_router.require_scoped_academic_enabled] = lambda: None
    app.dependency_overrides[staff_router.ManageDep.__metadata__[0].dependency] = lambda: staff_env[
        "admin"
    ]
    app.dependency_overrides[get_db] = lambda: staff_env["db"]
    with TestClient(app) as client:
        url = "/api/v1/colleges/1/staff/export-binding-tokens"
        for body in [
            {"staff_ids": []},
            {"days_valid": 91},
            {"staff_ids": ["1", "1"]},
            {"staff_ids": ["9223372036854775808"]},
            {"staff_ids": ["1"], "keyword": "教师"},
        ]:
            assert client.post(url, json=body).status_code == 422
        response = client.post(url, json={"days_valid": 30})
        assert response.status_code == 200
        assert response.headers["cache-control"] == "no-store"
        assert "attachment; filename*=UTF-8''" in response.headers["content-disposition"]
        assert response.headers["content-type"].endswith("spreadsheetml.sheet")
        assert load_workbook(BytesIO(response.content)).active["A2"].value == "00001"
