"""Unit tests for administrative class roster parsing."""

import io
import openpyxl
import pytest

from app.modules.importer.service import ImporterService


def _create_mock_school_roster() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "人工智能学院 数据科学2601 班级2026年课堂考勤表")
    ws.cell(2, 1, "序号")
    ws.cell(2, 2, "学号")
    ws.cell(2, 3, "姓名")
    ws.cell(4, 1, 1)
    ws.cell(4, 2, "0419260101")
    ws.cell(4, 3, "毕卓凡")
    ws.cell(5, 1, 2)
    ws.cell(5, 2, "0419260102")
    ws.cell(5, 3, "蔡景源")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _create_mock_standard_roster() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "学号")
    ws.cell(1, 2, "姓名")
    ws.cell(1, 3, "班级")
    ws.cell(1, 4, "学院")
    ws.cell(2, 1, "0419260201")
    ws.cell(2, 2, "张三")
    ws.cell(2, 3, "数据科学2602")
    ws.cell(2, 4, "人工智能学院")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_school_attendance_sheet():
    content = _create_mock_school_roster()
    service = ImporterService(None)  # unit test for parse only
    payload, issues, summary = service._parse_admin_roster(content, "数据科学2601.xlsx")

    assert summary["student_count"] == 2
    assert summary["classes_count"] == 1
    assert "数据科学2601" in summary["class_names"]
    assert len(payload["students"]) == 2
    s1 = payload["students"][0]
    assert s1["student_no"] == "0419260101"
    assert s1["name"] == "毕卓凡"
    assert s1["class_name"] == "数据科学2601"
    assert s1["college"] == "人工智能学院"
    assert s1["grade_year"] == 2026


def test_parse_standard_tabular_roster():
    content = _create_mock_standard_roster()
    service = ImporterService(None)
    payload, issues, summary = service._parse_admin_roster(content, "roster.xlsx")

    assert summary["student_count"] == 1
    assert summary["classes_count"] == 1
    assert "数据科学2602" in summary["class_names"]
    s = payload["students"][0]
    assert s["student_no"] == "0419260201"
    assert s["name"] == "张三"
    assert s["class_name"] == "数据科学2602"
    assert s["college"] == "人工智能学院"


def test_parse_class_duty_roster():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "二级学院学生责任区辅导员、班主任名单")
    ws.cell(2, 1, "序号")
    ws.cell(2, 2, "学院")
    ws.cell(2, 3, "班级")
    ws.cell(2, 4, "班主任")
    ws.cell(2, 5, "辅导员")
    ws.cell(3, 1, 13)
    ws.cell(3, 2, "人工智能学院")
    ws.cell(3, 3, "机电2501")
    ws.cell(3, 4, "黄朝阳")
    ws.cell(3, 5, "张一弛")
    buf = io.BytesIO()
    wb.save(buf)

    service = ImporterService(None)
    payload, issues, summary = service._parse_admin_roster(buf.getvalue(), "责任区名单.xlsx")
    assert len(payload["class_duties"]) == 1
    duty = payload["class_duties"][0]
    assert duty["class_name"] == "机电2501"
    assert duty["head_teacher"] == "黄朝阳"
    assert duty["counselor"] == "张一弛"
    assert duty["college"] == "人工智能学院"
    assert duty["grade_year"] == 2025


def test_parse_csv_duty_roster():
    csv_content = """二级学院学生责任区辅导员、班主任名单,,,,
序号,学院,班级,班主任,辅导员
13,人工智能学院,机电2501,黄朝阳,张一弛
15,人工智能学院,机自2621,黄建军,马敏敏
""".encode("utf-8")

    service = ImporterService(None)
    payload, issues, summary = service._parse_admin_roster(csv_content, "责任区名单.csv")
    assert len(payload["class_duties"]) == 2
    d1 = payload["class_duties"][0]
    assert d1["class_name"] == "机电2501"
    assert d1["head_teacher"] == "黄朝阳"
    assert d1["counselor"] == "张一弛"
    assert d1["grade_year"] == 2025

    d2 = payload["class_duties"][1]
    assert d2["class_name"] == "机自2621"
    assert d2["head_teacher"] == "黄建军"
    assert d2["counselor"] == "马敏敏"
    assert d2["grade_year"] == 2026

