from io import BytesIO

import pytest
from app.core.exceptions import AppError
from app.modules.academic.campus_models import CourseType
from app.modules.academic.models import AdministrativeClass, TeachingClass
from app.modules.importer.scoped_service import ScopedImporterService
from openpyxl import Workbook
from sqlalchemy import select

from tests.unit.test_campus_configuration import campus as campus_fixture
from tests.unit.test_campus_configuration import grant
from tests.unit.test_scoped_academic import records


@pytest.fixture(name="campus")
def configured():
    yield from campus_fixture.__wrapped__()


def timetable_sheet():
    book = Workbook()
    ws = book.active
    ws.append([
        "课程代码", "课程名称", "教学班码", "任课教师", "星期",
        "开始大节", "结束大节", "周次", "上课地点",
    ])
    ws.append([
        "CS101", "计算机导论", "CS101-01", "张老师", "1", "1", "2",
        "1,2,3,4,5,6,7,8", "教101",
    ])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def prepare_importer(campus):
    config, db, admin, teacher, a, b = campus
    _, _, _, teaching, _, own, foreign = records(db, a, b)
    grant(
        config, admin, teacher, a,
        ["import.execute", "course_schedule.import", "academic.manage"],
    )
    sid = teaching.semester_id
    from app.modules.academic.models import Semester
    from app.modules.academic.time_service import CollegeTimeService

    sem = db.get(Semester, sid)
    time_svc = CollegeTimeService(db, admin, int(a["id"]), sid)
    time_body = {
        "effective_from": sem.start_date.isoformat(),
        "periods": [
            {"period_no": p, "start_time": f"{7+p:02d}:00", "end_time": f"{7+p:02d}:45"}
            for p in range(1, 6)
        ],
        "grid_slot_mapping": {str(i): min(i, 5) for i in range(1, 13)},
    }
    time_svc.publish(time_body, time_svc.preview(time_body)["revision"])
    service = ScopedImporterService(db, teacher, int(a["id"]))
    return service, db, int(a["id"]), int(b["id"]), sid


def test_import_timetable_defaults_to_regular_course_type(campus):
    service, db, college_a, _, sid = prepare_importer(campus)
    content = timetable_sheet()

    preview = service.create_preview(
        service.actor,
        target="timetable",
        content=content,
        semester_id=sid,
        teaching_class_id=None,
        course_type_id=None,
        filename="timetable.xlsx",
        replace=False,
        request_id=None,
    )
    assert preview.can_confirm

    confirmed = service.confirm(service.actor, int(preview.id), request_id=None)
    assert confirmed.status == "CONFIRMED"

    regular = db.scalar(
        select(CourseType).where(
            CourseType.college_id == college_a,
            CourseType.builtin_key == "regular",
            CourseType.status == "ACTIVE",
        )
    )
    assert regular is not None

    tc = db.scalar(
        select(TeachingClass).where(
            TeachingClass.semester_id == sid,
            TeachingClass.class_code == "CS101-01",
        )
    )
    assert tc is not None
    assert tc.course_type_id == regular.id
    assert tc.college_id == college_a


def test_import_timetable_with_explicit_course_type(campus):
    service, db, college_a, _, sid = prepare_importer(campus)
    morning = db.scalar(
        select(CourseType).where(
            CourseType.college_id == college_a,
            CourseType.builtin_key == "morning_study",
            CourseType.status == "ACTIVE",
        )
    )
    assert morning is not None

    content = timetable_sheet()
    preview = service.create_preview(
        service.actor,
        target="timetable",
        content=content,
        semester_id=sid,
        teaching_class_id=None,
        course_type_id=morning.id,
        filename="morning_timetable.xlsx",
        replace=False,
        request_id=None,
    )
    assert preview.can_confirm

    confirmed = service.confirm(service.actor, int(preview.id), request_id=None)
    assert confirmed.status == "CONFIRMED"

    tc = db.scalar(
        select(TeachingClass).where(
            TeachingClass.semester_id == sid,
            TeachingClass.class_code == "CS101-01",
        )
    )
    assert tc is not None
    assert tc.course_type_id == morning.id


def test_import_timetable_rejects_foreign_course_type(campus):
    service, db, college_a, college_b, sid = prepare_importer(campus)
    foreign_type = db.scalar(
        select(CourseType).where(
            CourseType.college_id == college_b,
            CourseType.status == "ACTIVE",
        )
    )
    assert foreign_type is not None

    content = timetable_sheet()
    with pytest.raises(AppError) as exc_info:
        service.create_preview(
            service.actor,
            target="timetable",
            content=content,
            semester_id=sid,
            teaching_class_id=None,
            course_type_id=foreign_type.id,
            filename="foreign_timetable.xlsx",
            replace=False,
            request_id=None,
        )
    assert exc_info.value.http_status == 422
    assert exc_info.value.code == "VALIDATION_ERROR"


def test_bulk_import_with_course_type(campus):
    service, db, college_a, _, sid = prepare_importer(campus)
    ac = AdministrativeClass(
        college_id=college_a,
        class_code="\u8f6f\u5de52401",
        class_name="\u8f6f\u5de52401",
        status="ACTIVE",
    )
    db.add(ac)
    db.commit()

    book = Workbook()
    ws = book.active
    ws.title = "\u8f6f\u5de52401"
    ws.cell(row=1, column=1, value="\u8ba1\u7b97\u673a\u5b66\u9662\u73ed\u7ea7\u8bfe\u8868")
    ws.cell(row=2, column=1, value="\u73ed\u7ea7\uff1a\u8f6f\u5de52401")
    ws.cell(row=3, column=2, value="\u661f\u671f\u4e00")
    ws.cell(
        row=4,
        column=2,
        value=(
            "\u8f6f\u4ef6\u5de5\u7a0b\n"
            "\u25c71-8\u5468(1-2\u8282)\u25c7\u6559301"
            "\u25c7\u674e\u8001\u5e08\u25c7\u9009\u8bfe\u4eba\u6570\uff1a40"
        ),
    )
    stream = BytesIO()
    book.save(stream)
    grid_bytes = stream.getvalue()

    evening = db.scalar(
        select(CourseType).where(
            CourseType.college_id == college_a,
            CourseType.builtin_key == "evening_study",
            CourseType.status == "ACTIVE",
        )
    )
    assert evening is not None

    res = service.import_bulk_files(
        service.actor,
        semester_id=sid,
        files=[("\u8f6f\u5de52401\u8bfe\u8868.xlsx", grid_bytes)],
        course_type_id=evening.id,
    )
    assert res["total_files"] == 1
    assert res["file_results"][0]["status"] == "SUCCESS"

    tc = db.scalar(
        select(TeachingClass).where(
            TeachingClass.semester_id == sid,
            TeachingClass.college_id == college_a,
            TeachingClass.class_name.like("%\u8f6f\u4ef6\u5de5\u7a0b%"),
        )
    )
    assert tc is not None
    assert tc.course_type_id == evening.id
